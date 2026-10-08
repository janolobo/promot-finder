"""Public-source prospect research. No login automation or guessed personal data."""
from __future__ import annotations
import io
from html import unescape
import ipaddress
import json
import re
import select
import socket
import subprocess
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse, unquote
from urllib.robotparser import RobotFileParser

import requests
from catalogs import ANNOUNCEMENT_SOURCES, DIRECT_SOURCES, EUROPAGES_LEAD, collect, europages_search, load_internet_phrases, load_osm_phrases, match_analysis_phrases, match_csv_phrases, selected_announcement_phrases, selected_ares_phrases, selected_een_phrases, selected_europages_phrases, selected_hannovermesse_phrases, selected_industrystock_phrases, selected_wlw_phrases
# Announcements expire, so these catalogs are read again on every run instead of being marked complete.
REFRESHED_CATALOGS = ('techpilot', 'ariba', 'supplyon')
# Catalogs taken in full: the list itself is the selection, not the producer and recipient keywords.
WHOLE_CATALOGS = ('bvv', 'vdma', 'europages', 'wlw', 'industrystock', 'hannovermesse', 'ares') + ANNOUNCEMENT_SOURCES
ANALYSIS_GROUPS = {'competition': 'Konkurencja', 'cooperation': 'Współpraca', 'clients': 'Klienci'}
ANALYSIS_TITLES = {'competition': 'konkurencji', 'cooperation': 'współpracy', 'clients': 'klientów'}
ANALYSIS_MATCH_ORDER = ('Konkurencja', 'Współpraca', 'Klienci')
from local_osm import LocalIndex, GROUPS, qualify
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CATEGORIES = {
    'automotive': ('Podzespołów samochodowych', ['steering systems manufacturer', 'Hersteller Fahrwerk Lenkung', 'producent podzespołów samochodowych'], ['steering', 'suspension', 'automotive', 'lenkung', 'fahrwerk', 'zawiesze']),
    'agriculture': ('Maszyn rolniczych', ['agricultural machinery manufacturer', 'Landmaschinen Hersteller', 'producent maszyn rolniczych'], ['agricultur', 'landmaschinen', 'rolnicz', 'zeměděl']),
    'valves': ('Armatury i hydrauliki', ['industrial valve manufacturer', 'Industriearmaturen Hersteller', 'producent armatury hydraulicznej'], ['valve', 'armatur', 'hydraul', 'zawor', 'zawór']),
    'drives': ('Napędów i przekładni', ['gearbox coupling manufacturer', 'Getriebe Hersteller', 'producent przekładni napędów'], ['gearbox', 'coupling', 'getriebe', 'przekład', 'napęd']),
    'construction': ('Maszyn budowlanych i transportu', ['construction equipment manufacturer', 'Baumaschinen Hersteller', 'producent maszyn budowlanych'], ['construction equipment', 'baumaschinen', 'budowlan', 'transport']),
}
RECIPIENTS = {
    'forgings': ('Odkuwek', ['buyers of forgings', 'Abnehmer von Schmiedeteilen', 'odbiorcy odkuwek'], ['odkuw', 'schmiedeteil', 'forging']),
    'finished': ('Gotowych części', ['buyers of finished parts', 'Abnehmer von Fertigteilen', 'odbiorcy gotowych części'], ['gotowych części', 'gotowych czesci', 'części gotow', 'czesci gotow', 'fertigteil', 'finished part']),
    'iron': ('Odlewów żeliwnych', ['buyers of iron castings', 'Abnehmer von Eisengussteilen', 'odbiorcy odlewów żeliwnych'], ['żeliw', 'zeliw', 'gusseisen', 'iron casting']),
    'automotive': ('Automotive', ['automotive buyers', 'Automotive-Abnehmer', 'odbiorcy automotive'], ['automotive', 'motoryzac']),
    'forged': ('Elementów kutych', ['buyers of forged parts', 'Abnehmer von Schmiedeteilen', 'odbiorcy elementów kutych'], ['elementów kutych', 'elementow kutych', 'forged part', 'geschmiedet']),
}
PROFILES = {
    'all': ('Wszystkie rodzaje współpracy', ['', '', '']),
    'forgings': ('Odbiorcy odkuwek', ['forged components suppliers', 'Schmiedeteile Lieferanten', 'odkuwki dostawcy']),
    'finished_parts': ('Odbiorcy części gotowych', ['machined components suppliers', 'Fertigteile Zulieferer', 'części gotowe obrabiane dostawcy']),
    'cnc': ('Kooperacja CNC', ['CNC machining subcontracting', 'CNC Bearbeitung Kooperation', 'obróbka CNC kooperacja']),
}

def search_term(category, country, profile):
    language = {'DE': 1, 'PL': 2}.get(country, 0)
    return CATEGORIES[category][1][language] + (' ' + PROFILES[profile][1][language] if profile == 'cnc' else '')

def internet_query_terms(category_keys=None, recipient_keys=None, country='PL', profile='all'):
    """DuckDuckGo phrases from csv/internet_phrases.csv, without province names."""
    places = COUNTRY_QUERY_NAMES.get(country) or (COUNTRIES[country][0],)
    rows = load_internet_phrases()
    allowed = {key for key in (category_keys or []) + (recipient_keys or []) if key}
    if allowed:
        rows = [row for row in rows if row.get('category') in allowed]
    queries = []
    seen = set()
    for row in rows:
        phrase = ' '.join(str(row.get('phrase') or '').split())
        if not phrase:
            continue
        for place in places:
            query = phrase + ' ' + place
            key = query.casefold()
            if key not in seen:
                seen.add(key)
                queries.append(query)
    return queries

def recipient_term(key, country):
    language = {'DE': 1, 'PL': 2}.get(country, 0)
    return RECIPIENTS[key][1][language]

def audience_label(config):
    names = [RECIPIENTS[k][0] for k in config.get('recipients', []) if k in RECIPIENTS]
    return ', '.join(names) if names else PROFILES[config.get('profile', 'forgings')][0]

SOURCES = {
    'web': ('Internet — strony firm', ''),
    'fairs': ('Katalogi targowe', '(site:hannovermesse.de OR site:agritechnica.com OR site:automechanika.messefrankfurt.com)'),
    'tenders': ('Przetargi TED i BZP', '(site:ted.europa.eu OR site:ezamowienia.gov.pl)'),
    'directories': ('Katalogi branżowe i dystrybutorzy', '(distributor OR dealer OR dystrybutor OR Branchenverzeichnis)'),
    'pdf': ('Katalogi produktów PDF', 'filetype:pdf'),
}
COUNTRIES = {'PL': ('Polska', 'pl'), 'DE': ('Deutschland', 'de'), 'CZ': ('Česko', 'en')}
COUNTRY_QUERY_NAMES = {'PL': ('Polska', 'Poland'), 'DE': ('Deutschland', 'Germany'), 'CZ': ('Česko', 'Czechia', 'Czech Republic')}
WEB_FIRM_WORDS = {'PL': ('producent', 'firma'), 'DE': ('Hersteller', 'Unternehmen'), 'CZ': ('výrobce', 'firma')}
DENIED = ('linkedin.com', 'facebook.com', 'instagram.com', 'google.com', 'google.pl', 'google.de', 'maps.app.goo.gl', 'youtube.com')
DIRECTORIES = ('een.ec.europa.eu', 'hannovermesse.de', 'agritechnica.com', 'messefrankfurt.com', 'ted.europa.eu', 'ezamowienia.gov.pl', 'rejestr.io', 'gov.pl', 'ares.gov.cz', 'panoramafirm.pl', 'pkt.pl', 'europages.co.uk', 'europages.pl', 'europages.com', 'europages.de', 'kompass.com', 'wlw.de', 'industrystock.pl', 'industrystock.de', 'industrystock.com', 'lieferanten.de', 'sjn.de', 'factories.pl', 'pgm.org.pl', 'induux.de', 'ffo-info.de', 'yoys.pl', 'yoys.com')
EMAIL = re.compile(r'[A-Z0-9.!#$%&\'*+/=?^_`{|}~-]+@[A-Z0-9.-]+\.[A-Z]{2,}', re.I)
ROLE = re.compile(r'purchas|procurement|buyer|einkauf|zakup|sourcing|supplier|geschäftsführ|dyrektor|manager|sales|sprzedaż|vertrieb|director', re.I)
UA = 'PromotResearch/1.0 (public business directory research; respects robots.txt)'
OSM_SERVERS = {'privatecoffee': 'https://overpass.private.coffee/api/interpreter', 'fossgis': 'https://overpass-api.de/api/interpreter'}
COUNTRY_BBOX = {'DE': '47.2,5.8,55.1,15.1', 'CZ': '48.5,12,51.1,18.9'}
BUSINESS_FIELDS = ('nace', 'pkd', 'nip', 'regon', 'krs', 'ico', 'dic')

def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')

def domain(url):
    return (urlparse(url).hostname or '').lower().removeprefix('www.')

def is_domain(host, suffixes):
    return any(host == x or host.endswith('.' + x) for x in suffixes)

def safe_url(url):
    p = urlparse(url)
    if p.scheme not in ('http', 'https') or not p.hostname or p.username or p.password:
        raise ValueError('Niedozwolony adres URL')
    if p.port not in (None, 80, 443):
        raise ValueError('Dozwolone są standardowe porty stron WWW')
    addresses = socket.getaddrinfo(p.hostname, p.port or (443 if p.scheme == 'https' else 80))
    if not addresses or any(not ipaddress.ip_address(x[4][0]).is_global for x in addresses):
        raise ValueError('Adres lokalny lub prywatny — pominięto')
    return url

class Cancelled(Exception):
    pass

_LIVE_SESSIONS = set()
_LIVE_LOCK = threading.Lock()


def live_session(headers=None):
    """HTTP session that Stop can close from the panel thread."""
    session = requests.Session()
    if headers:
        session.headers.update(headers)
    with _LIVE_LOCK:
        _LIVE_SESSIONS.add(session)
    return session


def abort_network():
    """Drop in-flight HTTP so Stop does not wait for connect/read timeouts."""
    with _LIVE_LOCK:
        sessions = list(_LIVE_SESSIONS)
        _LIVE_SESSIONS.clear()
    for session in sessions:
        try:
            session.close()
        except Exception:
            pass


class Fetcher:
    def __init__(self, stop, log):
        self.stop, self.log = stop, log
        self.session = live_session({'User-Agent': UA})
        self.robots, self.last = {}, {}

    def check(self):
        if self.stop.is_set():
            raise Cancelled()

    def raw(self, url, **kwargs):
        for _ in range(6):
            self.check()
            safe_url(url)
            if is_domain(domain(url), DENIED):
                raise ValueError('Źródło dostępne tylko do ręcznego wyszukiwania')
            host = domain(url)
            wait = max(0, 1.1 - (time.monotonic() - self.last.get(host, 0)))
            if self.stop.wait(wait):
                raise Cancelled()
            self.last[host] = time.monotonic()
            try:
                response = self.session.get(url, timeout=(8, 18), allow_redirects=False, stream=True, **kwargs)
            except requests.RequestException:
                self.check()
                raise
            if response.is_redirect:
                target = urljoin(url, response.headers.get('Location', ''))
                response.close()
                if not self.allowed(target):
                    raise ValueError('Przekierowanie zablokowane przez robots.txt')
                url = target
                continue
            if response.status_code in (429, 503):
                delay = min(15, int(response.headers.get('Retry-After', '3')) if response.headers.get('Retry-After', '3').isdigit() else 3)
                response.close()
                if self.stop.wait(delay):
                    raise Cancelled()
                continue
            try:
                response.raise_for_status()
                chunks, size = [], 0
                for chunk in response.iter_content(65536):
                    self.check()
                    size += len(chunk)
                    if size > 6_000_000:
                        raise ValueError('Plik przekracza limit 6 MB')
                    chunks.append(chunk)
                return b''.join(chunks), response.headers.get('Content-Type', ''), url
            finally:
                response.close()
        raise ValueError('Przekroczony limit ponowień/przekierowań')

    def allowed(self, url):
        origin = '{0.scheme}://{0.netloc}'.format(urlparse(url))
        if origin not in self.robots:
            # Mark pending to avoid recursive robots redirects.
            self.robots[origin] = False
            try:
                data, _, _ = self.raw(origin + '/robots.txt')
                parser = RobotFileParser()
                parser.parse(data.decode('utf-8', 'replace').splitlines())
                self.robots[origin] = parser
            except requests.HTTPError as exc:
                self.robots[origin] = bool(exc.response is not None and exc.response.status_code == 404)
            except Cancelled:
                raise
            except Exception:
                self.robots[origin] = False
        rule = self.robots[origin]
        return rule if isinstance(rule, bool) else rule.can_fetch(UA, url)

    def page(self, url):
        if not self.allowed(url):
            raise ValueError('robots.txt nie pozwala na pobranie lub jest niedostępny')
        return self.raw(url)

def objects(value):
    if isinstance(value, dict):
        yield value
        for item in value.values():
            yield from objects(item)
    elif isinstance(value, list):
        for item in value:
            yield from objects(item)

def coords(lat, lon):
    try:
        a, b = float(lat), float(lon)
        return (a, b) if -90 <= a <= 90 and -180 <= b <= 180 else (None, None)
    except (ValueError, TypeError):
        return None, None

def extract_business_ids(text):
    """Read only explicitly labelled company identifiers and activity codes."""
    text = re.sub(r'\s+', ' ', str(text or '').replace('\xa0', ' '))
    result = {key: [] for key in BUSINESS_FIELDS}
    identifier_patterns = {
        'nip': (
            r'\bNIP\b(?:\s*(?:nr|number|no\.?))?\s*[:#-]?\s*(?:PL[\s-]*)?((?:\d[\s-]*){10})(?!\d)',
            r'\bVAT(?:\s*(?:ID|number|no\.?))?\s*[:#-]?\s*PL[\s-]*((?:\d[\s-]*){10})(?!\d)',
        ),
        'regon': (r'\bREGON\b(?:\s*(?:nr|number|no\.?))?\s*[:#-]?\s*((?:\d[\s-]*){14}|(?:\d[\s-]*){9})(?![\s-]*\d)',),
        'krs': (r'\bKRS\b(?:\s*(?:nr|number|no\.?))?\s*[:#-]?\s*((?:\d[\s-]*){10})(?!\d)',),
    }
    for key, patterns in identifier_patterns.items():
        for pattern in patterns:
            for match in re.finditer(pattern, text, re.I):
                value = re.sub(r'\D', '', match.group(1))
                if value and value not in result[key]:
                    result[key].append(value)
    labels = list(re.finditer(r'\b(PKD(?:\s+20\d{2})?|NACE(?:\s+Rev(?:ision)?\.?\s*\d+(?:\.\d+)?)?)\b', text, re.I))
    stop = re.compile(r'\b(?:NIP|REGON|KRS|VAT|address|adresse|adres|phone|telefon|e-?mail)\b', re.I)
    for index, label in enumerate(labels):
        key = 'pkd' if label.group(1).upper().startswith('PKD') else 'nace'
        end = min(len(text), label.end() + 240, labels[index + 1].start() if index + 1 < len(labels) else len(text))
        segment = stop.split(text[label.end():end], maxsplit=1)[0]
        pattern = r'(?<![\d.])\d{2}\.\d{2}\.[A-Z](?!\w)' if key == 'pkd' else r'(?<![\d.])(?:[A-U]\s*)?\d{2}\.\d{1,2}(?!\d|\.\d)'
        for value in re.findall(pattern, segment, re.I):
            value = re.sub(r'\s+', ' ', value.upper()).strip()
            if value not in result[key]:
                result[key].append(value)
    return result

def merge_business_ids(target, values):
    current = target.setdefault('business_ids', {})
    for key in BUSINESS_FIELDS:
        existing = current.setdefault(key, [])
        for value in values.get(key, []):
            if value not in existing:
                existing.append(value)

def extract(html, url):
    soup = BeautifulSoup(html, 'html.parser')
    result = {'name': '', 'address': '', 'lat': None, 'lon': None, 'geo_source': '', 'contacts': [], 'links': [], 'text': '', 'business_ids': {}}
    for script in soup.select('script[type="application/ld+json"]'):
        try:
            data = json.loads(script.string or script.get_text())
        except (ValueError, TypeError):
            continue
        for obj in objects(data):
            types = obj.get('@type', [])
            types = [types] if isinstance(types, str) else types
            if 'Person' in types and isinstance(obj.get('name'), str) and any(obj.get(k) for k in ['jobTitle', 'email', 'telephone']):
                emails = EMAIL.findall(str(obj.get('email', '')))
                phone = str(obj.get('telephone', '')) if isinstance(obj.get('telephone', ''), str) else ''
                result['contacts'].append(dict(person=obj['name'], role=str(obj.get('jobTitle', '')), email='; '.join(emails), phone=phone, source=url, status='Osoba — dane strukturalne strony, do weryfikacji'))
            if any(t in types for t in ['Organization', 'LocalBusiness', 'Corporation', 'Store', 'ProfessionalService']):
                email = '; '.join(EMAIL.findall(str(obj.get('email', ''))))
                phone = obj.get('telephone', '')
                if email or isinstance(phone, str) and phone:
                    result['contacts'].append(dict(person='', role='', email=email, phone=phone if isinstance(phone, str) else '', source=url, status='Kontakt firmy — dane strukturalne'))
            if 'latitude' in obj and 'longitude' in obj and result['lat'] is None:
                result['lat'], result['lon'] = coords(obj['latitude'], obj['longitude'])
                result['geo_source'] = url if result['lat'] is not None else ''
            if any(obj.get(k) for k in ('streetAddress', 'postalCode', 'addressLocality')) and not result['address']:
                result['address'] = ', '.join(str(obj.get(k, '')) for k in ['streetAddress', 'postalCode', 'addressLocality', 'addressCountry'] if isinstance(obj.get(k), str) and obj.get(k))
            if any(t in types for t in ['Organization', 'LocalBusiness', 'Corporation', 'AutomotiveBusiness', 'Store', 'ProfessionalService']) and not result['name']:
                result['name'] = str(obj.get('name', ''))
    from geolocation import address_rank, find_address
    for box in soup.select('footer, [itemprop=address], address'):
        found = find_address(box.get_text(' ', strip=True))
        if address_rank(found) > address_rank(result['address']):
            result['address'] = found
    street = soup.select_one('[itemprop=streetAddress]')
    postal = soup.select_one('[itemprop=postalCode]')
    locality = soup.select_one('[itemprop=addressLocality]')
    if street or postal or locality:
        joined = ', '.join(node.get_text(' ', strip=True) for node in (street, postal, locality) if node and node.get_text(strip=True))
        if address_rank(joined) > address_rank(result['address']):
            result['address'] = joined
    if result['lat'] is None:
        destinations = set(re.findall(r'https://www\.google\.com/maps/dir/[^\s\"\'<>]*?[?&]destination=(-?\d+(?:\.\d+)?),(-?\d+(?:\.\d+)?)', unescape(unquote(str(soup)))))
        if len(destinations) == 1:
            result['lat'], result['lon'] = coords(*next(iter(destinations)))
            result['geo_source'] = url if result['lat'] is not None else ''
    for node in soup(['script', 'style', 'noscript']):
        node.decompose()
    result['text'] = soup.get_text(' ', strip=True)[:100000]
    result['business_ids'] = extract_business_ids(result['text'])
    found = find_address(result['text'])
    if address_rank(found) > address_rank(result['address']):
        result['address'] = found
    if not result['name']:
        result['name'] = soup.title.get_text(' ', strip=True)[:180] if soup.title else domain(url)
    # Explicit Person microdata preserves association, unlike page-wide email matching.
    for card in soup.select('[itemtype$="/Person"]'):
        def field(key):
            tag = card.select_one('[itemprop="' + key + '"]')
            return (tag.get('content') or tag.get_text(' ', strip=True)) if tag else ''
        if field('name'):
            result['contacts'].append(dict(person=field('name'), role=field('jobTitle'), email='; '.join(EMAIL.findall(field('email'))), phone=field('telephone'), source=url, status='Osoba — microdata, do weryfikacji'))
    for a in soup.select('a[href]'):
        href = a.get('href', '').strip()
        if href.startswith('mailto:'):
            for email in EMAIL.findall(unquote(href.split('?')[0])):
                # A small contact card is evidence for a candidate name, never definitive attribution.
                parent = a.find_parent(['p', 'li', 'article', 'section', 'div'])
                context = parent.get_text(' ', strip=True) if parent else ''
                if len(context) > 500:
                    context = ''
                candidates = re.findall(r'\b([A-ZĄĆĘŁŃÓŚŹŻÄÖÜ][a-ząćęłńóśźżäöüß]+(?:-[A-ZĄĆĘŁŃÓŚŹŻÄÖÜ][a-ząćęłńóśźżäöüß]+)? [A-ZĄĆĘŁŃÓŚŹŻÄÖÜ][a-ząćęłńóśźżäöüß]+(?:-[A-ZĄĆĘŁŃÓŚŹŻÄÖÜ][a-ząćęłńóśźżäöüß]+)?)\b', context)
                title_words = {'inżynier', 'sprzedaży', 'kierownik', 'dyrektor', 'dział', 'zakupów', 'manager', 'director', 'sales', 'purchasing', 'contact', 'kontakt', 'general', 'customer', 'service', 'export', 'technical'}
                candidates = list(dict.fromkeys(n for n in candidates if not any(w.lower() in title_words for w in n.split())))
                person = candidates[0] if len(candidates) == 1 and ROLE.search(context) else ''
                result['contacts'].append(dict(person=person, role=context[:240] if person else '', email=email, phone='', source=url, status='Kandydat na osobę — sprawdź kontekst' if person else 'Kontakt strony — bez przypisania do osoby'))
        elif href.startswith('tel:'):
            phone = unquote(href[4:]).split('?')[0][:60]
            digits = re.sub(r'\D', '', phone)
            if 7 <= len(digits) <= 15 and digits not in ('0123456789', '123456789') and len(set(digits)) > 1:
                result['contacts'].append(dict(person='', role='', email='', phone=phone, source=url, status='Telefon strony — bez przypisania do osoby'))
        else:
            full = urljoin(url, href).split('#')[0]
            if full.startswith(('https://', 'http://')):
                result['links'].append((full, a.get_text(' ', strip=True)[:100]))
    for phone in re.findall(r'(?:telefon|telephone|phone|tel\.?)\s*[:.]?\s*(\+\d[\d\s()./\-]{6,}\d)', result['text'], re.I):
        digits = re.sub(r'\D', '', phone)
        if 7 <= len(digits) <= 15 and not any(re.sub(r'\D', '', item.get('phone') or '') == digits for item in result['contacts']):
            result['contacts'].append(dict(person='', role='', email='', phone=re.sub(r'\s+', ' ', phone).strip(), source=url, status='Telefon strony — bez przypisania do osoby'))
    existing = {c['email'].lower() for c in result['contacts']}
    for email in sorted(set(EMAIL.findall(result['text']))):
        if email.lower() not in existing:
            result['contacts'].append(dict(person='', role='', email=email, phone='', source=url, status='Kontakt w tekście — bez przypisania do osoby'))
    for contact in result['contacts']:
        if contact.get('email'):
            contact['email'] = '; '.join(dict.fromkeys(part.lower() for part in re.split(r'\s*;\s*', contact['email']) if part))
    return result

def company_list(html, url):
    """Read explicit business entries, never contacts from the whole directory."""
    soup = BeautifulSoup(html, 'html.parser')
    entries, seen = [], set()
    def append(name, target, fragment):
        if not isinstance(name, str) or not name.strip() or not isinstance(target, str):
            return
        target = urljoin(url, target)
        if not target or urlparse(target).scheme not in ('http', 'https') or target == url or target in seen:
            return
        seen.add(target)
        parsed = extract(fragment, target)
        fragment_soup = BeautifulSoup(fragment, 'html.parser')
        address_node = fragment_soup.select_one('[itemprop=address], address')
        if address_node and not parsed['address']:
            parsed['address'] = address_node.get_text(' ', strip=True)
        entries.append(dict(name=name.strip(), target=target, parsed=parsed))
    for script in soup.select('script[type="application/ld+json"]'):
        try:
            data = json.loads(script.string or script.get_text())
        except (ValueError, TypeError):
            continue
        for obj in objects(data):
            if obj.get('@type') != 'ItemList':
                continue
            for entry in obj.get('itemListElement', []):
                if not isinstance(entry, dict):
                    continue
                item = entry.get('item', entry)
                if isinstance(item, str):
                    item = dict(name=entry.get('name', ''), url=item)
                if isinstance(item, dict):
                    append(item.get('name', entry.get('name', '')), item.get('url', item.get('@id', '')),
                           '<script type="application/ld+json">' + json.dumps(item).replace('</', r'<\/') + '</script>')
    for card in soup.select('[itemtype$="/LocalBusiness"], [itemtype$="/Organization"], .company-card, .company-item, .business-card, .listing-item,  .search-result, .company, [id^="comp_"]'):
        anchor = card.select_one('[itemprop="name"] a[href], a[itemprop="url"][href], h2 a[href], h3 a[href], a.company-name[href]')
        if anchor is None:
            continue
        name_node = card.select_one('[itemprop="name"], h2, h3, .company-name')
        append((name_node or anchor).get_text(' ', strip=True), anchor.get('href', ''), str(card))
    return entries

def classify(text, selected):
    lowered = text.lower()
    hits = []
    for key in selected:
        matched = [w for w in CATEGORIES[key][2] if w in lowered]
        if matched:
            word = matched[0]
            at = lowered.find(word)
            hits.append((CATEGORIES[key][0], text[max(0, at - 70):at + 150]))
    return '; '.join(x[0] for x in hits), ' | '.join(x[1] for x in hits)[:1000]

def matches_selection(text, config):
    """Producers and recipients combine with AND. An empty list adds no constraint."""
    body = text or ''
    lowered = body.lower()
    cats = config.get('categories') or []
    recs = [key for key in config.get('recipients') or [] if key in RECIPIENTS]
    if cats and not classify(body, cats)[0]:
        return False
    if recs and not any(any(word in lowered for word in RECIPIENTS[key][2]) for key in recs):
        return False
    return True

class SearchBlocked(ValueError):
    """Explicit provider refusal; do not retry or rotate providers."""


def _search_web_body(query, country, config):
    page = max(1, int(config.get('web_page') or 1))
    count = max(10, min(50, int(config.get('web_max_results') or 20)))
    if config.get('engine', 'duckduckgo') == 'brave':
        session = live_session()
        try:
            response = session.get('https://api.search.brave.com/res/v1/web/search', params={'q': query, 'count': min(20, count), 'offset': (page - 1) * 10, 'country': country, 'search_lang': COUNTRIES[country][1]}, headers={'X-Subscription-Token': config['brave_key'], 'Accept': 'application/json'}, timeout=(8, 20))
            response.raise_for_status()
            return response.json().get('web', {}).get('results', [])
        finally:
            try:
                session.close()
            except Exception:
                pass
    from ddgs import DDGS
    region = {'PL': 'pl-pl', 'DE': 'de-de', 'CZ': 'cz-cs'}[country]
    from ddgs.exceptions import DDGSException, RatelimitException, TimeoutException
    try:
        results = DDGS(timeout=15).text(query, region=region, safesearch='moderate', max_results=count, page=page, backend='duckduckgo')
    except RatelimitException:
        raise SearchBlocked('DuckDuckGo ograniczył zapytania — zatrzymuję wyszukiwanie') from None
    except TimeoutException:
        raise ValueError('DuckDuckGo: przekroczono czas odpowiedzi dla tego zapytania') from None
    except DDGSException as exc:
        message = str(exc).strip().lower()
        if message == 'no results found.':
            return []
        if re.search(r'captcha|ratelimit|rate.limit|too many requests|\b(403|429)\b', message):
            raise SearchBlocked('DuckDuckGo odmówił dostępu lub ograniczył zapytania') from None
        raise ValueError('DuckDuckGo nie dostarczył poprawnej odpowiedzi dla tego zapytania') from None
    return [{'title': r.get('title', ''), 'url': r.get('href', r.get('url', ''))} for r in results]


def search_web(query, country, config):
    """One chosen provider, no proxy rotation or CAPTCHA bypass."""
    stop = config.get('stop')
    if isinstance(stop, threading.Event) and stop.is_set():
        raise Cancelled()
    if not isinstance(stop, threading.Event):
        return _search_web_body(query, country, config)
    box = {}
    def run():
        try:
            box['value'] = _search_web_body(query, country, config)
        except Exception as exc:
            box['error'] = exc
    thread = threading.Thread(target=run, daemon=True)
    thread.start()
    while thread.is_alive():
        if stop.is_set():
            raise Cancelled()
        thread.join(0.1)
    if 'error' in box:
        raise box['error']
    return box.get('value') or []


_CRAWL_LINK_GROUPS = (
    (0, r'contact|kontakt|impressum|imprint|anschrift|legal|mentions?.l[ée]gales|contatti|contacto|'
        r'iletisim|iletişim|adres|adresse|address|standort|location|lokaliz|dojazd|siedzib'),
    (1, r'team|people|staff|management|direction|equipe|squadra|ekibimiz|'
        r'einkauf|procurement|purchas|zakup|compras|acquisti|sales|sprzeda|verkauf|vertrieb|export'),
    (2, r'about|company|unternehmen|firma|o[-_ ]?nas|o[-_ ]?firmie|chi[-_ ]?siamo|'
        r'acerca|hakkimizda|hakkımızda|kurumsal'),
    (3, r'product|produkt|oferta|offer|service|leistung|branchen|industr'),
)
_CRAWL_LINK_SKIP = re.compile(
    r'privacy|datenschutz|cookie|terms|warunki|regulamin|blog|news|aktualno|career|karriere|'
    r'jobs?|login|register|cart|basket|checkout|facebook|instagram|linkedin|youtube',
    re.I,
)


def crawl_link_priority(url, label=''):
    """Priority of a public company subpage; lower values are read first."""
    text = unquote((urlparse(str(url or '')).path + ' ' + str(label or ''))).casefold()
    for priority, pattern in _CRAWL_LINK_GROUPS:
        if re.search(pattern, text, re.I):
            if priority > 0 and _CRAWL_LINK_SKIP.search(text):
                return None
            return priority
    return None


class Research:
    def __init__(self, folder):
        self.folder = Path(folder)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.stop = threading.Event()
        self.running = False
        self.records, self.discoveries, self.logs = [], [], []
        self.progress = dict(phase='Gotowy', done=0, total=0, queries_done=0, queries_total=0, errors=0)
        self.run_id = ''
        self.config = {}
        self.worker = None
        self.local_index = None
        self.local_indexes = {}
        # Live panel: wyniki/ sits next to cache/. Other folders keep a private cache.
        if self.folder.name == 'wyniki':
            self.cache_path = self.folder.parent / 'cache' / 'records.json'
        else:
            self.cache_path = self.folder / 'cache' / 'records.json'
        self.cache_path.parent.mkdir(parents=True, exist_ok=True)
        self.cached_records, self.cached_discoveries = [], []
        self.cache_counts = {}
        self.cache_history_migrated = False
        self.cache_last_write = 0.0
        try:
            cached = json.loads(self.cache_path.read_text(encoding='utf-8'))
            self.cached_records = cached.get('records', [])
            self.cached_discoveries = cached.get('discoveries', [])
            self.cache_counts = cached.get('catalog_counts', {})
            self.cache_history_migrated = bool(cached.get('history_migrated'))
            self.cache_counts = {key: 0 for key in (*DIRECT_SOURCES, *SOURCES, 'osm')}
            for item in self.cached_records:
                catalog_key = self._catalog_key(item)
                if catalog_key:
                    self.cache_counts[catalog_key] = self.cache_counts.get(catalog_key, 0) + 1
        except (OSError, ValueError, TypeError):
            pass
        files = sorted(self.folder.glob('*/wyniki.json'))
        if files and not self.cache_history_migrated:
            for history_file in files:
                try:
                    historical = json.loads(history_file.read_text(encoding='utf-8'))
                    self.cached_records.extend(historical.get('records', []))
                    self.cached_discoveries.extend(historical.get('discoveries', []))
                except (OSError, ValueError, TypeError):
                    continue
            self.cache_history_migrated = True
        if files:
            try:
                saved = json.loads(files[-1].read_text(encoding='utf-8'))
                self.records = saved.get('records', [])
                self.discoveries = saved.get('discoveries', [])
                self.logs = saved.get('logs', [])
                self.progress = saved.get('progress', self.progress)
                self.run_id = saved.get('run_id') or files[-1].parent.name
                self.config = saved.get('config', {})
            except (OSError, ValueError, TypeError):
                pass
        elif self.cached_records:
            self.records = json.loads(json.dumps(self.cached_records, ensure_ascii=False))
            self.discoveries = json.loads(json.dumps(self.cached_discoveries, ensure_ascii=False))
            self.progress.update(phase='Wczytano rekordy ze stałego cache', done=len(self.records), total=len(self.records))
        if self.records:
            self.backup_cache(force=True)

    def log(self, message):
        with self.lock:
            entry = f'{datetime.now():%H:%M:%S}  {message}'
            self.logs.append(entry)
            self.logs = self.logs[-2000:]
            if self.run_id:
                with (self.folder / self.run_id / 'poszukiwania.log').open('a', encoding='utf-8') as f:
                    f.write(entry + '\n')

    def snapshot(self):
        with self.lock:
            saved_config={key:self.config[key] for key in ('countries','categories','recipients','sources','pages','enrich_web','auto_translate') if key in self.config}
            return json.loads(json.dumps(dict(running=self.running, records=self.records, discoveries=self.discoveries, logs=self.logs, progress=self.progress, run_id=self.run_id, config=saved_config, cache_counts=self.cache_counts), ensure_ascii=False))

    @staticmethod
    def _cache_key(item, discovery=False):
        if discovery:
            return '\x1f'.join(str(item.get(key) or '').strip().casefold() for key in ('source','url','query','title'))
        source = str(item.get('source') or item.get('website') or '').strip().casefold()
        return '\x1f'.join((str(item.get('catalog') or '').strip().casefold(), source, str(item.get('catalog_name') or item.get('name') or '').strip().casefold()))

    def backup_cache(self, force=False):
        if not force and time.monotonic() - self.cache_last_write < 30:
            return
        with self.lock:
            records = {self._cache_key(item): item for item in self.cached_records if self._cache_key(item)}
            for item in self.records:
                key = self._cache_key(item)
                if not key:
                    continue
                old = records.get(key)
                if old and not item.get('checked_at'):
                    merged = dict(old)
                    for field, value in item.items():
                        if value not in ('', None, [], {}):
                            merged[field] = value
                    records[key] = merged
                else:
                    kept = dict(item)
                    if old and old.get('osm_text') and not kept.get('osm_text'):
                        kept['osm_text'] = old['osm_text']
                    records[key] = kept
            discoveries = {self._cache_key(item, True): item for item in self.cached_discoveries if self._cache_key(item, True)}
            discoveries.update({self._cache_key(item, True): item for item in self.discoveries if self._cache_key(item, True)})
            self.cached_records = list(records.values())
            self.cached_discoveries = list(discoveries.values())
            self.cache_counts = {key: 0 for key in (*DIRECT_SOURCES, *SOURCES, 'osm')}
            for item in self.cached_records:
                key = self._catalog_key(item)
                if key:
                    self.cache_counts[key] = self.cache_counts.get(key, 0) + 1
            payload = {
                'schema_version': 1,
                'updated_at': now(),
                'history_migrated': self.cache_history_migrated,
                'catalog_counts': self.cache_counts,
                'records': self.cached_records,
                'discoveries': self.cached_discoveries,
                'last_state': {
                    'run_id': self.run_id,
                    'config': self.config,
                    'progress': self.progress,
                    'logs': self.logs,
                },
            }
            temporary = self.cache_path.with_suffix('.tmp')
            temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
            temporary.replace(self.cache_path)
            self.cache_last_write = time.monotonic()

    @staticmethod
    def _catalog_key(item):
        labels = {label: key for key, (label, _url) in DIRECT_SOURCES.items()}
        labels.update({label: key for key, (label, _query) in SOURCES.items()})
        labels.update({'BVV / MSV — wystawcy': 'bvv', 'Europages — frazy produktowe': 'europages'})
        key = labels.get(str(item.get('catalog') or ''))
        if not key and 'openstreetmap.org' in str(item.get('source') or ''):
            return 'osm'
        return key or ''

    @staticmethod
    def catalog_name(key):
        if key == 'osm':
            return 'OSM'
        if key in DIRECT_SOURCES:
            return DIRECT_SOURCES[key][0]
        if key in SOURCES:
            return SOURCES[key][0]
        return key

    def clear_catalog(self, key):
        """Drop saved firms and leads from one direct catalog and allow that catalog to be read again."""
        if self.running:
            raise ValueError('Poczekaj na zakończenie bieżącego zadania')
        if key not in set(DIRECT_SOURCES) | set(SOURCES) | {'osm'}:
            raise ValueError('Nieznany katalog')
        label = self.catalog_name(key)
        aliases = {label}
        if key == 'bvv':
            aliases.add('BVV / MSV — wystawcy')
        if key == 'europages':
            aliases.add('Europages — frazy produktowe')
        with self.lock:
            before = {self._cache_key(item) or id(item) for item in list(self.cached_records) + list(self.records) if self._catalog_key(item) == key}
            self.records = [item for item in self.records if self._catalog_key(item) != key]
            self.cached_records = [item for item in self.cached_records if self._catalog_key(item) != key]
            self.discoveries = [item for item in self.discoveries if str(item.get('source') or '') not in aliases and not (key == 'osm' and 'openstreetmap.org' in str(item.get('url') or ''))]
            self.cached_discoveries = [item for item in self.cached_discoveries if str(item.get('source') or '') not in aliases and not (key == 'osm' and 'openstreetmap.org' in str(item.get('url') or ''))]
            resume = self.progress.setdefault('search_resume', {'completed_catalogs': [], 'phrases': [], 'queries': [], 'vdma_page': 1})
            resume['completed_catalogs'] = [item for item in (resume.get('completed_catalogs') or []) if item != key]
            phrases = list(resume.get('phrases') or [])
            if key == 'een':
                phrases = [item for item in phrases if not str(item).startswith('een:')]
            elif key == 'vdma':
                resume['vdma_page'] = 1
            elif key in ('europages', 'wlw', 'industrystock', 'hannovermesse', 'ares'):
                from catalogs import load_ares_phrases, load_europages_phrases, load_hannovermesse_phrases, load_industrystock_phrases, load_wlw_phrases
                loader = {'europages': load_europages_phrases, 'wlw': load_wlw_phrases, 'industrystock': load_industrystock_phrases, 'hannovermesse': load_hannovermesse_phrases, 'ares': load_ares_phrases}[key]
                owned = {row.get('phrase') for row in loader()}
                phrases = [item for item in phrases if item not in owned]
            resume['phrases'] = phrases
            if key in SOURCES:
                marker = SOURCES[key][1]
                tokens = []
                for item, spec in SOURCES.items():
                    if item == key or not spec[1]:
                        continue
                    tokens.extend(re.findall(r'site:[\w.-]+|filetype:\w+', spec[1]))
                    tokens.extend(re.findall(r'\b(?:distributor|dealer|dystrybutor|Branchenverzeichnis)\b', spec[1]))
                kept = []
                for query in resume.get('queries') or []:
                    text = str(query)
                    if marker:
                        if marker not in text:
                            kept.append(query)
                    elif any(token in text for token in tokens):
                        kept.append(query)
                resume['queries'] = kept
            removed = len(before)
        self.backup_cache(force=True)
        if self.run_id:
            path = self.folder / self.run_id / 'wyniki.json'
            temporary = path.with_suffix('.tmp')
            temporary.write_text(json.dumps(self.snapshot(), ensure_ascii=False, indent=2), encoding='utf-8')
            temporary.replace(path)
        self.log(f'Usunięto {removed} firm z katalogu {label}.')
        return removed

    def take_catalogs(self, config):
        """Catalog checkboxes chosen for fill, analysis or translation. None means every saved firm."""
        if not isinstance(config, dict) or 'catalogs' not in config:
            return None
        raw = config.get('catalogs')
        if not isinstance(raw, list) or any(not isinstance(item, str) for item in raw):
            raise ValueError('Nieprawidłowy wybór katalogów')
        allowed = set(DIRECT_SOURCES) | set(SOURCES) | {'osm'}
        catalogs = [item for item in raw if item in allowed]
        if not catalogs:
            raise ValueError('Zaznacz co najmniej jeden bezpośredni katalog')
        return catalogs

    def selected_records(self, catalogs, adopt=False):
        """Firms from the checked catalogs. The current session row wins over the saved copy."""
        wanted = set(catalogs or [])
        chosen = {}
        order = []
        for item in list(self.cached_records) + list(self.records):
            if self._catalog_key(item) not in wanted:
                continue
            key = self._cache_key(item) or id(item)
            if key not in chosen:
                order.append(key)
            chosen[key] = item
        rows = [chosen[key] for key in order]
        if adopt:
            present = {self._cache_key(item) for item in self.records if self._cache_key(item)}
            for item in rows:
                key = self._cache_key(item)
                if key and key not in present:
                    self.records.append(item)
                    present.add(key)
        return rows

    def catalog_rows(self):
        """Compact saved firms for the result list, grouped by direct catalog."""
        with self.lock:
            items = list(self.cached_records) + list(self.records)
        merged = {}
        for item in items:
            key = self._catalog_key(item)
            catalog = str(item.get('catalog') or '')
            if key == 'osm' and not catalog:
                catalog = 'OSM'
            if not key:
                continue
            ident = self._cache_key(item) or key + '\x1f' + str(item.get('name') or '')
            contacts = []
            for contact in (item.get('contacts') or [])[:8]:
                if contact.get('email') or contact.get('phone') or contact.get('person'):
                    contacts.append({field: contact.get(field, '') for field in ('person', 'role', 'email', 'phone', 'source', 'status')})
            merged[ident] = dict(
                id=f'{key}:{item.get("id") or ident}',
                name=item.get('name') or '',
                catalog=catalog,
                catalog_key=key,
                website=item.get('website') or '',
                source=item.get('source') or '',
                address=item.get('address') or '',
                lat=item.get('lat'),
                lon=item.get('lon'),
                province=item.get('province') or '',
                geo_precision=item.get('geo_precision') or '',
                location_scope=item.get('location_scope') or 'Lokalizacja niepotwierdzona',
                status=item.get('status') or 'Oczekuje',
                groups=item.get('groups') or ['Do weryfikacji'],
                contacts=contacts,
                business_ids=item.get('business_ids') or {},
                regon=item.get('regon') or {},
                establishments=item.get('establishments') or [],
                osm_text=str(item.get('osm_text') or '')[:500],
                description_pl=str(item.get('description_pl') or '')[:500],
                description_language=item.get('description_language') or '',
                published_at=str(item.get('published_at') or '')[:10],
            )
        return list(merged.values())

    def abort(self):
        """Stop from the panel thread: flag the worker and cut open HTTP."""
        self.stop.set()
        abort_network()
        process = getattr(self, 'job_process', None)
        if process is not None and process.poll() is None:
            try:
                process.terminate()
            except Exception:
                pass
        with self.lock:
            if self.running:
                self.progress['phase'] = 'Zatrzymywanie…'

    def checkpoint(self):
        if self.run_id:
            path = self.folder / self.run_id / 'wyniki.json'
            temp = path.with_suffix('.tmp')
            temp.write_text(json.dumps(self.snapshot(), ensure_ascii=False, indent=2), encoding='utf-8')
            temp.replace(path)
        if not self.stop.is_set():
            self.backup_cache()

    def start(self, config):
        with self.lock:
            if self.running:
                raise ValueError('Poszukiwania już trwają')
            self.backup_cache(force=True)
            profile = config.get('profile', 'forgings')
            if profile not in PROFILES:
                raise ValueError('Wybierz profil współpracy')
            cats = [k for k in config.get('categories', []) if k in CATEGORIES]
            raw_recipients = config.get('recipients', [])
            if not isinstance(raw_recipients, list) or any(not isinstance(item, str) for item in raw_recipients):
                raise ValueError('Wybierz odbiorców z listy')
            recs = [k for k in raw_recipients if k in RECIPIENTS]
            checked_countries = [k for k in config.get('countries', []) if k in COUNTRIES]
            countries = list(checked_countries)
            sources = [k for k in config.get('sources', []) if k in SOURCES or k in DIRECT_SOURCES or k in ['osm', 'linkedin', 'maps']]
            if not countries:
                raise ValueError('Wybierz kraj')
            listed = any(k in DIRECT_SOURCES or k == 'osm' for k in sources) or bool(config.get('seeds', '').strip())
            web_phrases = bool(load_internet_phrases()) if any(k in SOURCES for k in sources) else False
            if not listed and not web_phrases and not (cats or recs):
                raise ValueError('Wybierz Internet, OSM albo katalog. Frazy Internetu i OSM są w plikach CSV.')
            if not any(s in SOURCES or s in DIRECT_SOURCES or s == 'osm' for s in sources) and not config.get('seeds', '').strip():
                raise ValueError('Wybierz automatyczne źródło (np. OSM) lub podaj strony firm. LinkedIn i Google Maps są źródłami ręcznymi.')
            if config.get('engine', 'duckduckgo') not in ('duckduckgo', 'brave'):
                raise ValueError('Wybierz DuckDuckGo albo Brave API')
            if config.get('engine', 'duckduckgo') == 'brave' and any(s in SOURCES for s in sources) and not config.get('brave_key', '').strip():
                raise ValueError('Wybrane źródła wyszukiwarkowe wymagają klucza Brave Search API. Bez klucza wybierz OSM lub podaj strony firm.')
            bbox = [float(x.strip()) for x in config.get('bbox', '49,14,55,24').split(',')]
            if len(bbox) != 4 or not (-90 <= bbox[0] < bbox[2] <= 90 and -180 <= bbox[1] < bbox[3] <= 180):
                raise ValueError('Obszar: południe,zachód,północ,wschód')
            max_firms = int(config.get('max_firms', 50))
            pages = int(config.get('pages', 4))
            if not 1 <= max_firms <= 10000 or not 1 <= pages <= 20:
                raise ValueError('Limit: 1–10000 firm i 1–20 stron na firmę')
            if config.get('osm_server', 'privatecoffee') not in OSM_SERVERS:
                raise ValueError('Nieznany serwer OSM')
            self.config = dict(config, profile=profile, recipients=recs, categories=cats, countries=countries, een_countries=checked_countries, sources=sources, bbox=bbox, max_firms=max_firms, pages=pages)
            self.logs = []
            signature = dict(categories=cats, recipients=recs, countries=countries, sources=sources)
            resume = dict(self.progress.get('search_resume') or {})
            if resume.get('signature') != signature:
                resume = {'completed_catalogs': [], 'phrases': [], 'queries': [], 'vdma_page': 1, 'signature': signature}
            else:
                resume.setdefault('completed_catalogs', [])
                resume.setdefault('phrases', [])
                resume.setdefault('queries', [])
                resume.setdefault('vdma_page', 1)
            errors = self.progress.get('errors', 0)
            self.progress = dict(phase='Wyszukiwanie źródeł', done=0, total=0, queries_done=0, queries_total=0, errors=errors, search_resume=resume)
            self.run_id = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
            (self.folder / self.run_id).mkdir()
            self.stop.clear()
            self.running = True
            self.worker = threading.Thread(target=self.run, daemon=True)
            self.worker.start()

    def shutdown(self):
        """Stop cooperatively and let the worker finish its partial exports."""
        self.abort()
        if self.worker and self.worker.is_alive():
            self.worker.join()

    def error(self, label, exc):
        with self.lock:
            self.progress['errors'] += 1
        # Do not expose request headers, query keys or traceback in UI/log.
        detail = str(exc)[:180] if isinstance(exc, ValueError) else 'źródło niedostępne; kontynuuję'
        if isinstance(exc, requests.HTTPError) and exc.response is not None:
            detail = f'HTTP {exc.response.status_code} — źródło niedostępne; kontynuuję'
        self.log(f'{label}: {type(exc).__name__} — {detail}')

    def apply_location_scope(self, record):
        if record.get('lat') is None:
            record['location_scope'] = 'Lokalizacja niepotwierdzona'
            return
        codes = set(self.config.get('countries') or [])
        names = {COUNTRIES[code][0] for code in codes if code in COUNTRIES}
        country = str(record.get('country') or '')
        record['location_scope'] = 'Wybrany kraj' if country in codes or country in names else 'Poza wybranym obszarem'

    def ensure_local_indexes(self, countries):
        """Load every selected country that has a complete local PBF extract."""
        for code in countries:
            if code in self.local_indexes:
                continue
            try:
                index = LocalIndex(self.folder.parent, code)
            except ValueError:
                continue
            if not index.pbf.exists():
                continue
            index.ensure(self.stop, self.log)
            self.local_indexes[code] = index
        self.local_index = self.local_indexes.get('PL')

    def local_match(self, record):
        for index in self.local_indexes.values():
            match = index.find(record)
            if match:
                return match
        return None

    def locate_local_addresses(self, records):
        for index in self.local_indexes.values():
            index.locate_addresses(records, self.stop, self.log)

    def relevant_countries(self):
        selected=self.config.get('countries') or []
        if selected:return selected
        found={str(record.get('country') or '').upper() for record in self.records}
        return [code for code in COUNTRIES if code in found] or ['PL']

    def countries_for(self, records):
        """Selected countries plus the country of each firm being located. ARES stays in Czechia."""
        selected = [code for code in (self.config.get('countries') or []) if code in COUNTRIES]
        aliases = {
            'CZ': 'CZ', 'CZECHIA': 'CZ', 'ČESKO': 'CZ', 'CESKO': 'CZ', 'ČESKÁ REPUBLIKA': 'CZ', 'CESKA REPUBLIKA': 'CZ',
            'DE': 'DE', 'DEUTSCHLAND': 'DE', 'GERMANY': 'DE', 'NIEMCY': 'DE',
            'PL': 'PL', 'POLSKA': 'PL', 'POLAND': 'PL', 'POLSKO': 'PL',
        }
        extra = []
        for record in records or []:
            code = aliases.get(str(record.get('country') or '').strip().upper(), '')
            if not code and 'ares.gov.cz' in str(record.get('source') or ''):
                code = 'CZ'
            if code in COUNTRIES and code not in selected and code not in extra:
                extra.append(code)
        return selected + extra or ['PL']

    def place_ares_records(self, records):
        """Put ARES offices on the map from the Czech extract, then from the address text."""
        pending = [record for record in records if 'ares.gov.cz' in str(record.get('source') or '') and not isinstance(record.get('lat'), (int, float))]
        if not pending:
            return
        self.ensure_local_indexes(self.countries_for(pending))
        self.locate_local_addresses(pending)
        for record in pending:
            if self.stop.is_set():
                raise Cancelled()
            if not isinstance(record.get('lat'), (int, float)):
                self.place_record(record)

    def pin_after_enrich(self, record):
        """Put one firm on the map from the address just found on its page."""
        if self.stop.is_set():
            raise Cancelled()
        source = str(record.get('source') or '')
        from geolocation import pin_query
        if 'ares.gov.cz' not in source and pin_query(record)[0]:
            self.log('Lokalizuję: ' + record['name'])
            self.place_record(record)
        elif 'ares.gov.cz' in source:
            self.place_ares_records([record])
        if self.local_indexes and record.get('lat') is None:
            match = self.local_match(record)
            if match:
                with self.lock:
                    record.update(lat=match['lat'], lon=match['lon'], geo_source=match['source'], country=match.get('country', ''), province=match['province'], geo_precision=match['geo_precision'], osm_check='Dopasowano lokalny OSM')
                    record['address'] = record.get('address') or match['address']
        with self.lock:
            if record.get('lat') is not None:
                record.setdefault('geo_precision', 'Współrzędne ze źródła')
            self.apply_location_scope(record)

    def settle_found(self, fetch, records):
        """Fill websites, contacts and addresses, and pin each firm as soon as its page is read."""
        records = list(records)
        if not records:
            return
        geocode = self.config.get('geocode', True)
        if geocode:
            self.ensure_local_indexes(self.countries_for(records))
        if self.config.get('enrich_web', True):
            pending = [record for record in records if not self.data_complete(record)]
            skipped = len(records) - len(pending)
            if skipped:
                self.log(f'Komplet danych: {skipped} nowych firm ma nazwę, opis, adres, telefon i e-mail — pomijam je przy uzupełnianiu.')
            with self.lock:
                self.progress.update(phase='Uzupełniam: strony firm, adresy, telefony i e-maile', done=0, total=len(pending), fill_stage='enrich')
            self.log(f'Uzupełniam dane {len(pending)} znalezionych firm i od razu lokalizuję je na mapie.')
            for record in pending:
                fetch.check()
                self.log('Strona i kontakty: ' + record['name'])
                try:
                    self.enrich(record, fetch)
                    if geocode:
                        self.pin_after_enrich(record)
                except Cancelled:
                    raise
                except Exception as exc:
                    with self.lock:
                        record['status'] = 'Błąd pobrania — do sprawdzenia'
                    self.error(record['name'], exc)
                with self.lock:
                    record['checked_at'] = now()
                    if self.data_complete(record):
                        self.mark_data_complete(record)
                    self.progress['done'] += 1
                self.checkpoint()
        if self.config.get('auto_translate'):
            self.translate_descriptions(records)
        if not geocode:
            self.log('Lokalizacja na mapie jest wyłączona.')
            return
        self.locate_found(fetch, records)

    def locate_found(self, fetch, records):
        """Finish leftover pins from local OSM after each firm already got a page-based pin."""
        if not records:
            return
        self.ensure_local_indexes(self.countries_for(records))
        from catalogs import refresh_ares_addresses
        refresh_ares_addresses(records, fetch, self.log)
        self.locate_local_addresses(records)
        self.place_ares_records(records)
        for record in records:
            fetch.check()
            try:
                if record.get('lat') is None:
                    self.pin_after_enrich(record)
                else:
                    with self.lock:
                        self.apply_location_scope(record)
                with self.lock:
                    if self.data_complete(record):
                        self.mark_data_complete(record)
            except Cancelled:
                raise
            except Exception as exc:
                self.error(record['name'], exc)
            self.checkpoint()

    def run(self):
        c = self.config
        fetch = Fetcher(self.stop, self.log)
        queue = list(self.records)
        seen = set()
        for record in queue:
            website = record.get('website') or ''
            source = record.get('source') or ''
            for ident in (source if 'openstreetmap.org' in source else '', source + '|' + str(record.get('name') or '') if record.get('catalog') or not website else '', source.split('?')[0] if 'een.ec.europa.eu' in source else '', domain(website)):
                if ident:
                    seen.add(ident)
        batch_limit = c['max_firms']
        c['max_firms'] = len(queue) + batch_limit
        initial_count = len(queue)
        with self.lock:
            self.records = queue
        resume = self.progress.setdefault('search_resume', {'completed_catalogs': [], 'phrases': [], 'queries': [], 'vdma_page': 1})
        active_catalog = ''
        def add(url, name='', source='', lat=None, lon=None, address='', contacts=None, osm_id='', listing_id=''):
            if url and not url.startswith(('https://', 'http://')):
                url = 'https://' + url
            host = domain(url)
            ident = osm_id or listing_id or host
            if not ident or ident in seen or is_domain(host, DENIED) or is_domain(host, DIRECTORIES) or 'europages.' in host:
                return
            if len(queue) >= c['max_firms']:
                return
            # Keep separate OSM sites, merge generic domain seeds where possible.
            if not osm_id and not listing_id and any(domain(x['website']) == host for x in queue):
                return
            seen.add(ident)
            with self.lock:
                queue.append(dict(id=str(len(queue)+1), name=name or host, website=url, source=source or url, lat=lat, lon=lon, address=address, geo_source=source if lat is not None else '', contacts=contacts or [], province='', groups=['Do weryfikacji'], role_evidence='', profile=audience_label(c), category='', evidence='', status='Oczekuje', checked_at='', manual_links={}, catalog=active_catalog))
        inspected = set()
        def discover(url):
            if not url.startswith(('https://', 'http://')):
                url = 'https://' + url
            if url in inspected or is_domain(domain(url), DENIED):
                return
            if len(queue) >= c['max_firms']:
                return
            inspected.add(url)
            try:
                body, kind, final = fetch.page(url)
                entries = company_list(body, final) if 'html' in kind or not kind else []
                if entries:
                    self.log(f'Lista firm: {final} — znaleziono {len(entries)} wpisów.')
                    for entry in entries:
                        parsed, target = entry['parsed'], entry['target']
                        internal = domain(target) == domain(final) or is_domain(domain(target), DIRECTORIES)
                        website = '' if internal else target
                        if internal:
                            for href, label in parsed['links']:
                                if domain(href) != domain(final) and not is_domain(domain(href), DENIED + DIRECTORIES) and re.search(r'website|homepage|strona|www', label, re.I):
                                    website = href
                                    break
                        for contact in parsed['contacts']:
                            contact['source'] = final
                        add(website, name=entry['name'], source=target, address=parsed['address'], lat=parsed['lat'], lon=parsed['lon'], contacts=parsed['contacts'], listing_id=target if not website else '')
                    self.checkpoint()
                elif not is_domain(domain(final), DIRECTORIES):
                    add(final, source=url)
                else:
                    self.log('Katalog: nie rozpoznano listy firm — ' + final)
            except Cancelled:
                raise
            except Exception as exc:
                self.error('Odczyt strony ' + domain(url), exc)
                if not is_domain(domain(url), DIRECTORIES):
                    add(url, source=url)
        try:
            if 'osm' in c['sources']:
                self.progress['phase'] = 'Pobieranie danych z serwisu'
                self.ensure_local_indexes(c['countries'])
            self.log(f'Start. Zachowuję {len(queue)} wcześniejszych firm i dodaję kolejne, maksymalnie {batch_limit} w tym przebiegu.')
            catalog_keys=[k for k in c['sources'] if k in DIRECT_SOURCES]
            self.progress.update(catalogs_total=len(catalog_keys),catalogs_done=0,catalog_entries=len(queue))
            for key in catalog_keys:
                fetch.check()
                if key in resume.get('completed_catalogs', []):
                    self.log(DIRECT_SOURCES[key][0] + ': katalog był już odczytany, zachowuję firmy i przechodzę dalej.')
                    self.progress['catalogs_done'] += 1
                    continue
                name,url=DIRECT_SOURCES[key]
                lead=dict(title=name,url=url,source=name,category='Katalog bezpośredni',query='',status='Odczyt katalogu',checked_at=now())
                self.discoveries.append(lead)
                imported=0;found=0;filtered=0
                try:
                    search = (lambda query: europages_search(query, c)) if key in ('europages', 'wlw', 'industrystock', 'hannovermesse') else None
                    europages_phrases = selected_europages_phrases(c.get('categories')) if key == 'europages' else None
                    wlw_phrases = selected_wlw_phrases(c.get('categories')) if key == 'wlw' else None
                    industrystock_phrases = selected_industrystock_phrases(c.get('categories')) if key == 'industrystock' else None
                    hannovermesse_phrases = selected_hannovermesse_phrases(c.get('categories')) if key == 'hannovermesse' else None
                    ares_phrases = selected_ares_phrases(c.get('categories')) if key == 'ares' else None
                    een_phrases = selected_een_phrases() if key == 'een' else None
                    announcement_phrases = selected_announcement_phrases(key) if key in ANNOUNCEMENT_SOURCES and key != 'een' else None
                    if key == 'europages':
                        self.log(f'Europages: plik csv/europages_100_phrases.csv — {len(europages_phrases)} fraz po filtrze kategorii.')
                    if key == 'wlw':
                        self.log(f'WLW: plik csv/wlw_100_phrases.csv — {len(wlw_phrases)} fraz po filtrze kategorii.')
                    if key == 'industrystock':
                        self.log(f'IndustryStock: plik csv/industrystock_phrases.csv — {len(industrystock_phrases)} fraz.')
                    if key == 'hannovermesse':
                        self.log(f'Hannover Messe: plik csv/hannovermesse_phrases.csv — {len(hannovermesse_phrases)} fraz.')
                    if key == 'ares':
                        self.log(f'ARES API: plik csv/ares_phrases.csv — {len(ares_phrases)} fraz po filtrze kategorii.')
                    if key == 'een':
                        self.log(f'Enterprise Europe Network: plik csv/een_phrases.csv — {len(een_phrases)} fraz.')
                    if announcement_phrases is not None:
                        self.log(f'{name}: plik csv/{key}_phrases.csv — {len(announcement_phrases)} haseł.')
                    done_phrases = set(resume.get('phrases') or [])
                    if key == 'europages' and europages_phrases is not None:
                        europages_phrases = [row for row in europages_phrases if row.get('phrase') not in done_phrases]
                    if key == 'wlw' and wlw_phrases is not None:
                        wlw_phrases = [row for row in wlw_phrases if row.get('phrase') not in done_phrases]
                    if key == 'industrystock' and industrystock_phrases is not None:
                        industrystock_phrases = [row for row in industrystock_phrases if row.get('phrase') not in done_phrases]
                    if key == 'hannovermesse' and hannovermesse_phrases is not None:
                        hannovermesse_phrases = [row for row in hannovermesse_phrases if row.get('phrase') not in done_phrases]
                    if key == 'ares' and ares_phrases is not None:
                        ares_phrases = [row for row in ares_phrases if row.get('phrase') not in done_phrases]
                    een_countries = [code for code in (c.get('een_countries') if 'een_countries' in c else c.get('countries') or []) if code in ('PL', 'DE', 'CZ')]
                    een_scope = ','.join(een_countries)
                    if key == 'een' and een_phrases is not None:
                        een_phrases = [row for row in een_phrases if f'een:{een_scope}:{row.get("phrase", "")}' not in done_phrases]
                    catalog_finished = not (key == 'een' and not een_countries)
                    current_phrase = ''
                    phrase_resume_ok = True
                    phrase_done = True
                    for item in collect(key, fetch, self.log, search=search, wlw_phrases=wlw_phrases, europages_phrases=europages_phrases, industrystock_phrases=industrystock_phrases, hannovermesse_phrases=hannovermesse_phrases, ares_phrases=ares_phrases, een_phrases=een_phrases, een_countries=een_countries, announcement_phrases=announcement_phrases, vdma_page=int(resume.get('vdma_page') or 1) if key == 'vdma' else 1):
                        phrase = item.get('query') or ''
                        if phrase and phrase != current_phrase:
                            if current_phrase and phrase_resume_ok:
                                resume.setdefault('phrases', [])
                                if current_phrase not in resume['phrases']:
                                    resume['phrases'].append(current_phrase)
                            current_phrase = phrase
                            phrase_resume_ok = item.get('resume_ok', True)
                        if item.get('resume_ok') is False:
                            phrase_resume_ok = False
                            catalog_finished = False
                        if item.get('een_skipped'):
                            continue
                        if item.get('vdma_page'):
                            resume['vdma_page'] = item['vdma_page']
                        found+=1
                        if item.get('lead_only'):
                            self.discoveries.append(dict(title=item['name'], url=item['source'], source=name, category=item.get('lead_category') or 'Profil Europages', query=item.get('query', ''), status=item.get('lead_status') or EUROPAGES_LEAD, checked_at=now()))
                            self.progress['catalog_entries'] += 1
                            self.checkpoint()
                            continue
                        country=item.get('country','').strip().casefold()
                        province=item.get('province','')
                        # MSV category 11 and the VDMA member list are the requested catalogs, including firms seated abroad.
                        abroad=key not in WHOLE_CATALOGS and country and country not in ('polska','poland','pl','polsko') and not c.get('foreign_catalogs',False)
                        outside=abroad
                        narrow=False if key in WHOLE_CATALOGS else not matches_selection(item.get('text',''), c)
                        excluded=outside or narrow
                        if outside:
                            lead_status='Poza wybranym obszarem — pominięto import'
                        elif narrow:
                            lead_status='Nie spełnia warunku oraz dla producentów i odbiorców'
                        else:
                            lead_status='Firma z publicznego indeksu ' + {'europages': 'Europages', 'wlw': 'WLW', 'industrystock': 'IndustryStock', 'hannovermesse': 'Hannover Messe'}.get(key, '') if key in ('europages', 'wlw', 'industrystock', 'hannovermesse') else 'Wpis firmy w katalogu'
                        self.discoveries.append(dict(title=item.get('announcement') or item['name'],url=item['source'],source=name,category=country or 'Kraj do sprawdzenia',query=item.get('query',''),status=lead_status,checked_at=now()))
                        if excluded:
                            filtered+=1
                            continue
                        before=len(queue)
                        listing_id = item['source'] if key in ANNOUNCEMENT_SOURCES else (item['source'] + '|' + item['name'] if key in ('bvv', 'vdma', 'ares') or not item['website'] else '')
                        add(item['website'],item['name'],item['source'],item.get('lat'),item.get('lon'),item['address'],item['contacts'],listing_id=listing_id)
                        if len(queue)>before:
                            row=queue[-1]
                            row.update(catalog=name,country=item.get('country',''),catalog_province=province,province=province,osm_text=item.get('text',''),groups=['Do weryfikacji'],role_evidence='',locality=item.get('locality',''),address_code=item.get('address_code',''),published_at=item.get('published_at') or '')
                            if item.get('business_ids'):
                                merge_business_ids(row, item['business_ids'])
                            if item.get('establishments'):
                                row['establishments'] = item['establishments']
                            if key in ('europages', 'wlw', 'industrystock', 'hannovermesse', 'ares') + ANNOUNCEMENT_SOURCES:
                                row['catalog_name'] = item['name']
                            imported+=1
                        self.progress['catalog_entries']+=1
                        self.checkpoint()
                        if len(queue)>=c['max_firms']:
                            catalog_finished = False
                            phrase_done = False
                            self.log('Limit firm — import katalogu częściowy. Następne uruchomienie doda kolejne firmy.');break
                    if current_phrase and phrase_done and phrase_resume_ok:
                        resume.setdefault('phrases', [])
                        if current_phrase not in resume['phrases']:
                            resume['phrases'].append(current_phrase)
                    if catalog_finished and key not in REFRESHED_CATALOGS:
                        resume.setdefault('completed_catalogs', [])
                        if key not in resume['completed_catalogs']:
                            resume['completed_catalogs'].append(key)
                        if key == 'vdma':
                            resume['vdma_page'] = 1
                    lead['status']=f'Odczytano {found}; dodano {imported}; poza obszarem {filtered}.'
                except Cancelled:
                    raise
                except Exception as exc:
                    lead['status']='Nie udało się ukończyć odczytu — sprawdź log'
                    self.error(name,exc)
                self.progress['catalogs_done']+=1
                self.log(name+': '+lead['status'])
                self.checkpoint()
            for url in c.get('seeds', '').splitlines():
                if url.strip():
                    discover(url.strip())
            def osm_service(box, label):
                fetch.check()
                self.log(f'OSM: wyszukiwanie w serwisie — {label}.')
                query = f'[out:json][timeout:35];(nwr["man_made"="works"]["name"]({box});nwr["industrial"]["name"]({box});nwr["craft"="metal_construction"]["name"]({box}););out center tags {c["max_firms"]};'
                endpoint = OSM_SERVERS[c.get('osm_server', 'privatecoffee')]
                self.log('Serwer OSM: ' + domain(endpoint))
                try:
                    response = fetch.session.get(endpoint, params={'data': query}, headers={'User-Agent': UA, 'Accept': 'application/json'}, timeout=(8, 45))
                except requests.RequestException:
                    fetch.check()
                    raise
                response.raise_for_status()
                payload = response.json()
                if payload.get('remark'):
                    self.log('OSM: serwer zgłosił niepełny wynik: ' + payload['remark'][:180])
                osm_limit = max(1, (c['max_firms'] - len(queue)) // 2) if any(s in SOURCES for s in c['sources']) else c['max_firms']
                osm_added = 0
                for item in payload.get('elements', []):
                    if osm_added >= osm_limit or len(queue) >= c['max_firms']:
                        break
                    t = item.get('tags', {})
                    haystack=' '.join(str(t.get(key) or '') for key in ('name','description','product','industrial','craft','operator'))
                    if c.get('categories') or c.get('recipients'):
                        if not matches_selection(haystack, c):
                            continue
                    elif not match_csv_phrases('osm', haystack):
                        continue
                    centre = item.get('center', item)
                    lat, lon = coords(centre.get('lat'), centre.get('lon'))
                    source = f'https://www.openstreetmap.org/{item["type"]}/{item["id"]}'
                    contact = dict(person='', role='', email=str(t.get('contact:email', t.get('email', '')) or '').lower(), phone=t.get('contact:phone', t.get('phone', '')), source=source, status='Kontakt OSM — do weryfikacji')
                    before = len(queue)
                    add(t.get('website', t.get('contact:website', '')), t.get('name', ''), source, lat, lon, ', '.join(t.get(k, '') for k in ['addr:street', 'addr:housenumber', 'addr:postcode', 'addr:city'] if t.get(k)), [contact] if contact['email'] or contact['phone'] else [], source)
                    if len(queue) > before:
                        osm_added += 1
                self.checkpoint()
            if 'osm' in c['sources']:
                self.log(f'OSM: plik csv/osm_phrases.csv — {len(load_osm_phrases())} fraz.')
                for code in c['countries']:
                    index=self.local_indexes.get(code)
                    if index is not None:
                        local=index.data['companies']
                        local.sort(key=lambda r: (not bool(classify(r['text'], c['categories'])[0]), qualify(r['text'])[0] == ['Do weryfikacji'], r['name']))
                        self.log(f'Lokalny OSM: {len(local)} obiektów w kraju {COUNTRIES[code][0]}; limit importu {c["max_firms"]}.')
                        ordered=iter(local)
                        for item in ordered:
                            fetch.check()
                            haystack=' '.join(str(item.get(key) or '') for key in ('name','text','address'))
                            if c.get('categories') or c.get('recipients'):
                                if not matches_selection(haystack, c): continue
                            elif not match_csv_phrases('osm', haystack):
                                continue
                            before=len(queue)
                            ct=dict(person='',role='',email=str(item['email'] or '').lower(),phone=item['phone'],source=item['source'],status='Kontakt obiektu OSM — do sprawdzenia')
                            add(item['website'],item['name'],item['source'],item['lat'],item['lon'],item['address'],[ct] if item['email'] or item['phone'] else [],item['source'])
                            if len(queue)>before:
                                queue[-1].update(country=code,province=item.get('province',''),geo_precision=item['geo_precision'],groups=['Do weryfikacji'],role_evidence='',osm_text=item['text'])
                            if len(queue)>=c['max_firms']: break
                        self.checkpoint()
                        continue
                    if code not in COUNTRY_BBOX and code!='PL':
                        continue
                    try:
                        box=','.join(map(str,c['bbox'])) if code=='PL' else COUNTRY_BBOX[code]
                        osm_service(box, COUNTRIES[code][0])
                    except Exception as exc:
                        self.error('OSM ' + COUNTRIES[code][0], exc)
            web = [s for s in c['sources'] if s in SOURCES]
            producer_keys = c['categories']
            recipient_keys = c.get('recipients', [])
            jobs = []
            if web:
                web_rows = load_internet_phrases()
                self.log(f'Internet: plik csv/internet_phrases.csv — {len(web_rows)} fraz.')
                for source in web:
                    for country in c['countries']:
                        for term in internet_query_terms(producer_keys, recipient_keys, country, c.get('profile', 'all')):
                            jobs.append((source, country, term, 'frazy CSV'))
            with self.lock:
                self.progress['queries_total'] = len(jobs)
            if 'web' in web and not jobs:
                self.log('Internet: brak fraz w csv/internet_phrases.csv — nie ma zapytania do DuckDuckGo.')
            done_queries = set(resume.get('queries') or [])
            pending = []
            skipped = 0
            for job in jobs:
                query = ' '.join(part for part in (job[2], COUNTRIES[job[1]][0], SOURCES[job[0]][1]) if part)
                if job[0] != 'web' and query in done_queries:
                    skipped += 1
                    continue
                pending.append(job)
            jobs = pending
            with self.lock:
                self.progress['queries_total'] = len(jobs) + skipped
                self.progress['queries_done'] = skipped
            if skipped:
                self.log(f'Pomijam {skipped} już wykonanych zapytań. Internet szuka od nowa.')
            search_failures = 0
            for source, country, term, label in jobs:
                fetch.check()
                if len(queue) >= c['max_firms']:
                    self.progress['search_stop_reason'] = 'Osiągnięto limit firm'
                    self.log(f'Internet: zebrano {len(queue) - initial_count} nowych firm — osiągnięto limit {batch_limit}.')
                    break
                query = term if source == 'web' else ' '.join(part for part in (term, COUNTRIES[country][0], SOURCES[source][1]) if part)
                if source != 'web':
                    query = ' '.join(part for part in (query, SOURCES[source][1]) if part and SOURCES[source][1] not in query)
                pages = 8 if source == 'web' else 1
                page_empty = 0
                for page in range(1, pages + 1):
                    if len(queue) >= c['max_firms']:
                        break
                    self.log('Szukam: ' + query + (f' · strona {page}' if page > 1 else ''))
                    try:
                        if self.stop.wait(0.2):
                            raise Cancelled()
                        remaining = max(10, c['max_firms'] - len(queue))
                        hits = search_web(query, country, dict(c, web_page=page, web_max_results=min(50, remaining), stop=self.stop))
                        search_failures = 0
                        self.log(f'Wyszukiwarka zwróciła {len(hits)} wyników.')
                        if not hits:
                            self.progress['queries_empty'] = self.progress.get('queries_empty', 0) + 1
                            self.log('Brak wyników zwróconych dla tego zapytania — przechodzę do kolejnego; to nie jest potwierdzenie braku firm.')
                            page_empty += 1
                            if page_empty >= 2:
                                break
                            continue
                        page_empty = 0
                        active_catalog = SOURCES[source][0]
                        added = 0
                        for hit in hits:
                            if len(queue) >= c['max_firms']:
                                break
                            url = hit.get('url') or hit.get('href') or ''
                            title = str(hit.get('title') or '').strip()
                            with self.lock:
                                if not any(d['url'] == url for d in self.discoveries):
                                    self.discoveries.append(dict(title=title, url=url, source=SOURCES[source][0], category=label, query=query, status='Trop — niepotwierdzona firma / potrzeba', checked_at=now()))
                            before = len(queue)
                            add(url, title, url)
                            if len(queue) > before:
                                queue[-1]['osm_text'] = title
                                queue[-1]['country'] = COUNTRIES[country][0]
                                added += 1
                                self.log('Na listę: ' + queue[-1]['name'])
                        self.log(f'Dodano {added} nazw firm z DuckDuckGo do listy. Internet: {len(queue) - initial_count}/{batch_limit} nowych w tym przebiegu.')
                        active_catalog = ''
                        if added == 0 and source == 'web':
                            break
                    except Cancelled:
                        raise
                    except Exception as exc:
                        self.error('Wyszukiwarka', exc)
                        search_failures += 1
                        self.progress['queries_failed'] = self.progress.get('queries_failed', 0) + 1
                        response = getattr(exc, 'response', None)
                        if isinstance(exc, SearchBlocked) or (response is not None and response.status_code in (401, 403, 429)):
                            self.progress['search_stop_reason'] = 'Dostawca odmówił dostępu lub ograniczył zapytania'
                            self.progress['queries_done'] += 1
                            self.log(self.progress['search_stop_reason'] + ' — zapisuję zebrane wyniki.')
                            page_empty = pages
                            break
                        if search_failures >= 5:
                            self.progress['search_stop_reason'] = 'Pięć kolejnych awarii wyszukiwarki — przerwano plan'
                            self.log(self.progress['search_stop_reason'])
                            page_empty = pages
                            break
                        self.log('Przerwa 5 sekund po błędzie; następnie kolejne zapytanie.')
                        if self.stop.wait(5):
                            raise Cancelled()
                        break
                with self.lock:
                    self.progress['queries_done'] += 1
                    resume.setdefault('queries', [])
                    if query not in resume['queries']:
                        resume['queries'].append(query)
                self.checkpoint()
                if self.progress.get('search_stop_reason') in ('Dostawca odmówił dostępu lub ograniczył zapytania', 'Pięć kolejnych awarii wyszukiwarki — przerwano plan'):
                    break
            with self.lock:
                self.progress['queries_skipped'] = self.progress['queries_total'] - self.progress['queries_done']
                self.records = queue
                self.progress.update(phase='Zebrano firmy — uzupełnij dane, potem analizuj', done=0, total=0)
            new_records = queue[initial_count:]
            self.log(f'Znaleziono {len(queue)} kandydatów, w tym {len(new_records)} nowych. Uzupełniam dane i od razu lokalizuję na mapie.')
            self.settle_found(fetch, new_records)
            with self.lock:
                if not queue:
                    self.progress['phase'] = 'Nie udało się wyszukać firm — błąd źródła; sprawdź log' if self.progress['errors'] else 'Brak firm dla tych ustawień — sprawdź tropy i zmień zapytania'
                elif c.get('enrich_web', True) and c.get('geocode', True):
                    self.progress['phase'] = 'Zakończono: zebrano firmy, uzupełniono dane i położono je na mapie'
                elif c.get('enrich_web', True):
                    self.progress['phase'] = 'Zakończono uzupełnianie danych — lokalizacja na mapie wyłączona'
                else:
                    self.progress['phase'] = 'Zakończono zbieranie — strony firm uzupełnisz przyciskiem Uzupełnij dane'
            if 'linkedin' in c['sources'] or 'maps' in c['sources']:
                self.log('LinkedIn / Google Maps: linki wyszukiwania ręcznego są przy firmach; serwisy nie były scrapowane.')
        except (Cancelled, InterruptedError):
            with self.lock:
                self.progress['phase'] = 'Zatrzymano — wyniki częściowe'
                for record in self.records:
                    if record['status'] in ('Analiza', 'Oczekuje'):
                        record['status'] = 'Zatrzymano — analiza nieukończona'
            self.log('Zatrzymano na żądanie użytkownika.')
        except Exception as exc:
            with self.lock:
                self.progress['phase'] = 'Błąd — wyniki częściowe'
            self.error('Poszukiwania', exc)
        finally:
            with self.lock:
                if not self.records:
                    self.records = queue
            self.log(f'Zapisuję {len(self.records)} firm i {len(self.discoveries)} tropów.')
            try:
                export_xlsx(self.snapshot(), self.folder / self.run_id / 'wyniki.xlsx')
                payload = json.dumps(self.snapshot(), ensure_ascii=False).replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026')
                template = (Path(__file__).parent / 'static' / 'map.html').read_text(encoding='utf-8')
                (self.folder / self.run_id / 'mapa.html').write_text(template.replace('/*SNAPSHOT*/null', payload), encoding='utf-8')
                self.log('Gotowe: wyniki.xlsx, mapa.html i log sesji.')
            except Exception as exc:
                self.error('Eksport plików', exc)
                with self.lock:
                    self.progress['phase'] += ' — błąd eksportu; sprawdź log'
            with self.lock:
                self.running = False
                self.checkpoint()

    def locate_records(self):
        fetch = Fetcher(self.stop, self.log)
        self.ensure_local_indexes(self.countries_for(self.records))
        from catalogs import refresh_ares_addresses
        refresh_ares_addresses(self.records, fetch, self.log)
        self.locate_local_addresses(self.records)
        with self.lock:
            self.progress.update(phase='Lokalizowanie firm na mapie', geo_done=0, geo_total=len(self.records))
        for record in self.records:
            if self.stop.is_set():
                raise Cancelled()
            if record.get('lat') is None and not record.get('catalog') and is_domain(domain(record.get('source', '')), DIRECTORIES):
                try:
                    body, kind, final = fetch.page(record['source'])
                    parsed = extract(body, final) if 'html' in kind else {}
                    with self.lock:
                        if parsed.get('address'):
                            record['address'] = parsed['address']
                        if parsed.get('lat') is not None:
                            record.update(lat=parsed['lat'], lon=parsed['lon'], geo_source=final, geo_precision='Współrzędne z wpisu firmy — do weryfikacji')
                except Cancelled:
                    raise
                except Exception as exc:
                    self.error('Adres z wpisu katalogowego', exc)
            from geolocation import pin_query
            if pin_query(record)[0]:
                self.log('Lokalizuję: ' + record['name'])
                self.place_record(record)
            if self.local_indexes and record.get('lat') is None:
                record['osm_check'] = 'Sprawdzono lokalny OSM — brak jednoznacznego dopasowania'
                match = self.local_match(record)
                if match:
                    record['osm_check'] = 'Dopasowano lokalny OSM'
                    record.update(lat=match['lat'],lon=match['lon'],geo_source=match['source'],country=match.get('country',''),province=match['province'],geo_precision=match['geo_precision'])
                    record['address'] = record.get('address') or match['address']
            with self.lock:
                if record.get('lat') is not None:
                    record.setdefault('geo_precision', 'Współrzędne ze źródła')
                self.apply_location_scope(record)
                self.progress['geo_done'] += 1
            if self.progress['geo_done'] % 20 == 0 or self.progress['geo_done'] == len(self.records):
                self.checkpoint()

    @staticmethod
    def visible_description(record):
        """Description shown in the list, without the company name and NACE/PKD codes."""
        def clean(text):
            firm = re.sub(r'\s+', ' ', str(record.get('name') or '')).strip().casefold()
            kept = []
            for part in re.split(r'\s*·\s*', str(text or '')):
                part = re.sub(r'\s+', ' ', part).strip()
                if not part or re.match(r'^(?:cz[-\s]*)?(?:nace|pkd)\b', part, re.I):
                    continue
                if firm and part.casefold() == firm:
                    continue
                kept.append(part)
            joined = re.sub(r'(?:^|[\s·])(?:cz[-\s]*)?(?:nace|pkd)\s*:\s*\d[\d\s,./-]*', ' ', ' · '.join(kept), flags=re.I)
            return re.sub(r'^[·,\s]+|[·,\s]+$', '', re.sub(r'\s+', ' ', joined)).strip()
        translated = clean(record.get('description_pl'))
        original = clean(re.sub(r'\s+', ' ', str(record.get('osm_text') or '')).strip())
        return translated or original

    @staticmethod
    def data_complete(record):
        """Name, description, address, phone and e-mail are all present."""
        contacts = record.get('contacts') or []
        phone = any(str(item.get('phone') or '').strip() for item in contacts)
        email = any(str(item.get('email') or '').strip() for item in contacts)
        return bool(str(record.get('name') or '').strip() and Research.visible_description(record) and str(record.get('address') or '').strip() and phone and email)

    @staticmethod
    def mark_data_complete(record):
        record['data_complete'] = True
        record['status'] = 'Komplet danych'

    def start_fill(self, config=None):
        self._launch(self.fill_worker, config)

    def start_analysis(self, config=None):
        self._launch(self.judgement_worker, config)

    def _launch(self, worker, config=None):
        config = config or {}
        with self.lock:
            if self.running:
                raise ValueError('Poczekaj na zakończenie bieżącego zadania')
            filling = worker == self.fill_worker
            if not self.records and not (filling and self.europages_pending()):
                files = sorted(self.folder.glob('*/wyniki.json'))
                if not files:
                    raise ValueError('Najpierw wyszukaj firmy')
                saved = json.loads(files[-1].read_text(encoding='utf-8'))
                self.records = saved.get('records', [])
                self.discoveries = saved.get('discoveries', [])
                self.progress = saved.get('progress', self.progress)
                self.run_id = files[-1].parent.name
            if not self.records and not (filling and self.europages_pending()):
                raise ValueError('Najpierw wyszukaj firmy')
            if 'pages' in config:
                pages = int(config['pages'])
                if not 1 <= pages <= 20:
                    raise ValueError('Limit: 1–20 stron na firmę')
                self.config['pages'] = pages
            elif 'pages' not in self.config:
                self.config['pages'] = 4
            if 'enrich_web' in config:
                self.config['enrich_web'] = bool(config['enrich_web'])
            if 'auto_translate' in config:
                self.config['auto_translate'] = bool(config['auto_translate'])
            if isinstance(config.get('categories'), list):
                self.config['categories'] = [k for k in config['categories'] if k in CATEGORIES]
            kind = str(config.get('analysis_kind') or '').strip()
            if kind and kind not in ANALYSIS_GROUPS and kind != 'all':
                raise ValueError('Nieznany rodzaj analizy')
            if worker == self.judgement_worker:
                self.config['analysis_kind'] = kind or 'all'
            catalogs = self.take_catalogs(config)
            if catalogs is None:
                self.config.pop('job_catalogs', None)
            else:
                self.config['job_catalogs'] = catalogs
                if not self.selected_records(catalogs) and not (filling and 'europages' in catalogs and self.europages_pending()):
                    raise ValueError('W zaznaczonych katalogach nie ma firm')
            self.logs = []
            self.stop.clear()
            self.running = True
            self.worker = threading.Thread(target=worker, daemon=True)
            self.worker.start()

    def europages_pending(self):
        return [item for item in self.discoveries if item.get('status') == EUROPAGES_LEAD]

    def promote_europages_leads(self):
        """Turn a saved name and profile URL into a firm row. The profile page is not the data source."""
        known = {record.get('source') for record in self.records}
        added = 0
        for item in self.discoveries:
            if item.get('status') != EUROPAGES_LEAD or item.get('url') in known:
                continue
            with self.lock:
                self.records.append(dict(id=str(len(self.records) + 1), name=item.get('title') or '', website='', source=item.get('url') or '', lat=None, lon=None, address='', geo_source='', contacts=[], province='', groups=['Do weryfikacji'], role_evidence='', profile='', category='', evidence='', status='Oczekuje', checked_at='', manual_links={}, catalog=item.get('source') or '', osm_text='', country=''))
                item['status'] = 'Przekazany do uzupełnienia — strona firmy, rejestr, OSM'
            known.add(item.get('url'))
            added += 1
        if added:
            self.log(f'Europages: {added} tropów idzie do uzupełnienia. Adres i kontakty biorę ze strony firmy, a punkt na mapie z tego adresu. REGON zostaje w widoku PKD.')
        return added

    def fill_worker(self):
        fetch = Fetcher(self.stop, self.log)
        try:
            catalogs = self.config.get('job_catalogs')
            if catalogs is None or 'europages' in catalogs:
                self.promote_europages_leads()
            all_records = list(self.records) if catalogs is None else self.selected_records(catalogs, adopt=True)
            if catalogs is not None and not all_records:
                raise ValueError('W zaznaczonych katalogach nie ma firm')
            previous_phase = str(self.progress.get('phase') or '')
            previous_done = int(self.progress.get('done') or 0)
            previous_total = int(self.progress.get('total') or 0)
            previous_offset = int(self.progress.get('fill_offset') or 0)
            previous_stage = str(self.progress.get('fill_stage') or '')
            resume_from = 0
            if previous_phase.startswith('Zatrzymano uzupełnianie') and (catalogs is None or list(self.progress.get('fill_catalogs') or []) == list(catalogs)):
                if previous_stage == 'locate':
                    resume_from = len(all_records)
                elif 0 < previous_total <= len(all_records):
                    inferred_offset = previous_offset or len(all_records) - previous_total
                    if previous_total == len(all_records) - inferred_offset:
                        resume_from = min(len(all_records), inferred_offset + previous_done)
            original_names={item.get('url'):item.get('title') for item in self.discoveries if 'ARES' in str(item.get('source') or '') and item.get('url') and item.get('title')}
            from catalogs import _ARES_SITE_SKIP
            retry_records = []
            complete_ids = set()
            for record in all_records:
                if self.data_complete(record):
                    self.mark_data_complete(record)
                    complete_ids.add(id(record))
                    continue
                if domain(record.get('source') or '')=='ares.gov.cz':
                    official=record.get('catalog_name') or original_names.get(record.get('source'))
                    if official:
                        record['name']=record['catalog_name']=official
                    if record.get('website') and (record.get('ares_site_verified')!=3 or is_domain(domain(record['website']),_ARES_SITE_SKIP)):
                        record.update(website='',contacts=[],page_text='',status='Oczekuje')
                        retry_records.append(record)
            if complete_ids:
                self.log(f'Komplet danych: {len(complete_ids)} firm ma nazwę, opis, adres, telefon i e-mail — pomijam je przy uzupełnianiu.')
                self.checkpoint()
            records = retry_records + [record for record in all_records[resume_from:] if record not in retry_records and id(record) not in complete_ids]
            if resume_from:
                self.log(f'Wznawiam uzupełnianie od firmy {resume_from + 1}; wcześniej sprawdzonych firm nie pobieram ponownie.')
            with self.lock:
                self.progress.update(phase='Uzupełniam: strony firm, adresy, telefony i e-maile', done=0, total=len(records), fill_stage='enrich', fill_offset=resume_from, fill_catalogs=list(catalogs or []))
            scope = 'ze wszystkich katalogów' if catalogs is None else 'z katalogów: ' + ', '.join(self.catalog_name(key) for key in catalogs)
            self.log(f'Szukam stron, adresów i kontaktów {len(records)} firm {scope} i od razu lokalizuję je na mapie.')
            self.ensure_local_indexes(self.countries_for(records or all_records))
            for record in records:
                fetch.check()
                if 'europages.' in domain(record.get('source') or '') and not record.get('catalog_name'):
                    from catalogs import europages_company
                    found = europages_company('', record.get('source') or '')
                    if found:
                        with self.lock:
                            record['catalog_name'] = found[0]
                            record['name'] = found[0]
                self.log('Strona i kontakty: ' + record['name'])
                try:
                    if self.config.get('enrich_web', True):
                        self.enrich(record, fetch)
                    self.pin_after_enrich(record)
                    source_host = domain(record.get('source') or '')
                    with self.lock:
                        if ('europages.' in source_host or source_host in ('wlw.de', 'industrystock.com', 'industrystock.de', 'industrystock.pl', 'hannovermesse.de')) and record.get('status') in ('Oczekuje', 'Zatrzymano — analiza nieukończona', 'Nie udało się odczytać strony', 'Błąd pobrania — do sprawdzenia'):
                            useful = bool(record.get('address') or record.get('contacts') or record.get('page_text'))
                            record['status'] = 'Dane uzupełnione — oczekuje na analizę' if useful else 'Brak danych na stronie — do sprawdzenia'
                except Cancelled:
                    raise
                except Exception as exc:
                    with self.lock:
                        record['status'] = 'Błąd pobrania — do sprawdzenia'
                    self.error(record['name'], exc)
                with self.lock:
                    record['checked_at'] = now()
                    if self.data_complete(record):
                        self.mark_data_complete(record)
                    self.progress['done'] += 1
                self.checkpoint()
            if self.config.get('auto_translate'):
                self.translate_descriptions(None if catalogs is None else all_records)
            locate_rows = [record for record in all_records if id(record) not in complete_ids]
            from catalogs import refresh_ares_addresses
            refresh_ares_addresses(locate_rows, fetch, self.log)
            self.locate_local_addresses(locate_rows)
            for record in locate_rows:
                fetch.check()
                try:
                    if record.get('lat') is None:
                        self.pin_after_enrich(record)
                    source_host = domain(record.get('source') or '')
                    with self.lock:
                        self.apply_location_scope(record)
                        if ('europages.' in source_host or source_host in ('wlw.de', 'industrystock.com', 'industrystock.de', 'industrystock.pl', 'hannovermesse.de')) and record.get('status') in ('Oczekuje', 'Zatrzymano — analiza nieukończona', 'Nie udało się odczytać strony', 'Błąd pobrania — do sprawdzenia'):
                            useful = bool(record.get('address') or record.get('contacts') or record.get('page_text'))
                            record['status'] = 'Dane uzupełnione — oczekuje na analizę' if useful else 'Brak danych na stronie — do sprawdzenia'
                        if self.data_complete(record):
                            self.mark_data_complete(record)
                except Cancelled:
                    raise
                except Exception as exc:
                    with self.lock:
                        record['status'] = 'Błąd pobrania — do sprawdzenia'
                    self.error(record['name'], exc)
                self.checkpoint()
            with self.lock:
                self.progress['phase'] = 'Zakończono uzupełnianie danych'
            self.log('Uzupełnianie danych zakończone.')
        except Cancelled:
            with self.lock:
                self.progress['phase'] = 'Zatrzymano uzupełnianie — wyniki częściowe'
            self.log('Zatrzymano uzupełnianie na żądanie użytkownika.')
        except Exception as exc:
            with self.lock:
                self.progress['phase'] = 'Błąd uzupełniania — wyniki częściowe'
            self.error('Uzupełnianie danych', exc)
        finally:
            self._finish_job('Eksport po uzupełnieniu')

    def judgement_worker(self):
        try:
            catalogs = self.config.get('job_catalogs')
            records = list(self.records) if catalogs is None else self.selected_records(catalogs, adopt=True)
            if catalogs is not None and not records:
                raise ValueError('W zaznaczonych katalogach nie ma firm')
            scope = 'ze wszystkich katalogów' if catalogs is None else 'z katalogów: ' + ', '.join(self.catalog_name(key) for key in catalogs)
            with self.lock:
                self.progress.update(phase='Analiza firm', done=0, total=len(records))
            kind = self.config.get('analysis_kind') or 'all'
            title = 'wszystkie analizy' if kind == 'all' else ANALYSIS_TITLES.get(kind, 'firm')
            self.log(f'Analizuję {len(records)} firm {scope}: {title}, słowa z CSV i status.')
            for record in records:
                if self.stop.is_set():
                    raise Cancelled()
                self.log('Analizuję: ' + record['name'])
                with self.lock:
                    record['status'] = 'Analiza'
                    self.apply_judgement(record)
                    record['checked_at'] = now()
                self.place_record(record)
                with self.lock:
                    self.progress['done'] += 1
                self.checkpoint()
            with self.lock:
                self.progress['phase'] = 'Zakończono analizę firm'
            self.log('Analiza firm zakończona.')
        except Cancelled:
            with self.lock:
                self.progress['phase'] = 'Zatrzymano analizę — wyniki częściowe'
                for record in self.records:
                    if record['status'] == 'Analiza':
                        record['status'] = 'Zatrzymano — analiza nieukończona'
            self.log('Zatrzymano analizę na żądanie użytkownika.')
        except Exception as exc:
            with self.lock:
                self.progress['phase'] = 'Błąd analizy — wyniki częściowe'
            self.error('Analiza firm', exc)
        finally:
            self._finish_job('Eksport po analizie')

    def _finish_job(self, export_label):
        self.checkpoint()
        self.backup_cache(force=True)
        try:
            if self.run_id:
                export_xlsx(self.snapshot(), self.folder / self.run_id / 'wyniki.xlsx')
        except Exception as exc:
            self.error(export_label, exc)
        with self.lock:
            self.running = False
            self.checkpoint()

    def analysis_kinds(self):
        kind = self.config.get('analysis_kind') or 'all'
        if kind in ANALYSIS_GROUPS:
            return [kind]
        return list(ANALYSIS_GROUPS)

    def apply_judgement(self, record):
        """Assign legend groups from one analysis CSV or from all three."""
        body = ' '.join(part for part in (record.get('name', ''), record.get('osm_text', ''), record.get('page_text', ''), record.get('description_pl', ''), record.get('address', '')) if part)
        kinds = self.analysis_kinds()
        labels = {ANALYSIS_GROUPS[kind] for kind in kinds}
        groups = [item for item in record.get('groups') or [] if item not in labels and item != 'Do weryfikacji']
        hits = []
        for kind in kinds:
            label = ANALYSIS_GROUPS[kind]
            found = match_analysis_phrases(kind, body)
            if found:
                groups.append(label)
                hits.extend(found)
        record['groups'] = groups or ['Do weryfikacji']
        record['analysis_match'] = ' · '.join(item for item in ANALYSIS_MATCH_ORDER if item in record['groups'])
        if hits:
            record['role_evidence'] = 'Słowa z CSV analizy: ' + ', '.join(hits[:8])
            record['category'] = record['analysis_match']
            record['evidence'] = ' | '.join(hits)[:3000]
            record['status'] = 'Dopasowanie słów — wymaga kwalifikacji'
        else:
            record['role_evidence'] = 'Brak słów z CSV analizy'
            record['category'] = ''
            record['evidence'] = ''
            record['status'] = 'Do sprawdzenia'

    def place_record(self, record):
        """Put a checked firm on the map from its address, or from the city when that is all we have."""
        from geolocation import Locator, city_note, pin_query
        query, city_only = pin_query(record)
        address = record.get('address') or record.get('locality') or ''
        has_point = isinstance(record.get('lat'), (int, float))
        city_pin = str(record.get('geo_precision') or '').startswith('Miasto')
        if has_point and not (query and not city_only and city_pin):
            return
        if getattr(self, 'analysis_locator', None) is None:
            self.analysis_locator = Locator(self.folder, self.stop)
        if not str(query).strip():
            return
        result = self.analysis_locator.locate(query)
        if not isinstance(result, dict) or result.get('lat') is None:
            return
        if city_only or city_note(address):
            result = dict(result, geo_precision='Miasto — przybliżenie, nie siedziba firmy')
        with self.lock:
            record.update(lat=result['lat'], lon=result['lon'], geo_source=result.get('geo_source', ''), geo_precision=result.get('geo_precision', ''), geo_label=result.get('geo_label', ''))

    def translate_descriptions(self, records=None):
        """Translate German and Czech descriptions locally through isolated Argos."""
        records = self.records if records is None else records
        worker = Path(__file__).parent / 'argos_worker.py'
        python = Path(__file__).parent / '.venv-argos' / 'bin' / 'python'
        if not python.exists():
            self.log('Argos Translate: brak lokalnego środowiska .venv-argos — pomijam tłumaczenie.')
            self.progress['phase'] = 'Brak lokalnego tłumacza Argos'
            return 0
        pending = []
        for record in records:
            text = re.sub(r'\s+', ' ', str(record.get('osm_text') or '')).strip()[:4000]
            source = self.description_language(record, text)
            if text and source and record.get('description_source') != text:
                pending.append(dict(id=str(len(pending)), text=text, source=source, record=record))
        if not pending:
            self.log('Argos Translate: brak nowych niemieckich lub czeskich opisów.')
            return 0
        with self.lock:
            self.progress['phase'] = 'Tłumaczenie opisów — Argos Translate'
        self.log(f'Argos Translate: tłumaczę lokalnie {len(pending)} opisów.')
        process = subprocess.Popen(
            [str(python), str(worker)],
            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            text=True, encoding='utf-8',
        )
        self.job_process = process
        payload = json.dumps([{key: item[key] for key in ('id', 'text', 'source')} for item in pending], ensure_ascii=False)
        jobs = {job['id']: job for job in pending}
        saved = 0
        seen = 0
        total = len(pending)
        try:
            process.stdin.write(payload)
            process.stdin.close()
            while True:
                if self.stop.is_set():
                    process.terminate()
                    process.wait(timeout=5)
                    raise Cancelled()
                ready = True
                fileno = getattr(process.stdout, 'fileno', None)
                if callable(fileno):
                    try:
                        ready = bool(select.select([process.stdout], [], [], 0.5)[0])
                    except (OSError, TypeError, ValueError):
                        ready = True
                if not ready:
                    if process.poll() is not None:
                        break
                    continue
                line = process.stdout.readline()
                if not line:
                    break
                item = json.loads(line)
                seen += 1
                job = jobs.get(str(item.get('id')))
                name = str((job or {}).get('record', {}).get('name') or 'opis')
                translation = str(item.get('translation') or '').strip()
                if job and translation:
                    with self.lock:
                        job['record']['description_pl'] = translation
                        job['record']['description_source'] = item.get('text', '')
                        job['record']['description_language'] = item.get('source', '')
                        self.progress['phase'] = f'Tłumaczenie opisów: {seen}/{total}'
                    saved += 1
                    self.log(f'Argos Translate: {seen}/{total} — {name}')
                    if saved % 25 == 0:
                        with self.lock:
                            self._copy_translations_to_cache()
                        self.checkpoint()
                else:
                    with self.lock:
                        self.progress['phase'] = f'Tłumaczenie opisów: {seen}/{total}'
                    self.log(f'Argos Translate: błąd tłumaczenia {seen}/{total} — {name}')
            stderr = process.stderr.read() if process.stderr else ''
            process.wait(timeout=5)
            if process.returncode:
                detail = (stderr or '').strip().splitlines()[-1:] or ['nieznany błąd']
                self.log('Argos Translate: nie udało się przetłumaczyć — ' + detail[0][:180])
                if not saved:
                    self.progress['phase'] = 'Nie udało się przetłumaczyć opisów'
                    return 0
            with self.lock:
                self._copy_translations_to_cache()
            self.log(f'Argos Translate: zapisano {saved} tłumaczeń.')
            self.checkpoint()
            return saved
        except (OSError, ValueError, subprocess.SubprocessError) as exc:
            if process.poll() is None:
                process.terminate()
            self.log(f'Argos Translate: błąd lokalnego tłumacza — {type(exc).__name__}.')
            return saved
        finally:
            if getattr(self, 'job_process', None) is process:
                self.job_process = None

    def _copy_translations_to_cache(self):
        live = {self._cache_key(item): item for item in self.records if self._cache_key(item)}
        for item in self.cached_records:
            current = live.get(self._cache_key(item))
            if not current:
                continue
            if current.get('osm_text'):
                item['osm_text'] = current['osm_text']
            elif item.get('osm_text') and not current.get('osm_text'):
                current['osm_text'] = item['osm_text']
            for field in ('description_pl', 'description_source', 'description_language'):
                if current.get(field):
                    item[field] = current[field]

    def firms_to_translate(self):
        """Saved catalog firms plus the current session. The session row is the one updated in the list."""
        chosen = {}
        for item in self.cached_records:
            key = self._cache_key(item)
            if key:
                chosen[key] = item
        for item in self.records:
            key = self._cache_key(item) or id(item)
            chosen[key] = item
        return list(chosen.values())

    def start_translate(self, config=None):
        catalogs = self.take_catalogs(config or {})
        with self.lock:
            if self.running:
                raise ValueError('Poczekaj na zakończenie bieżącego zadania')
            if not self.records and not self.cached_records:
                raise ValueError('Najpierw wyszukaj firmy')
            if catalogs is not None and not self.selected_records(catalogs):
                raise ValueError('W zaznaczonych katalogach nie ma firm')
            self._translate_catalogs = catalogs
            self.stop.clear()
            self.running = True
            self.progress['phase'] = 'Tłumaczenie opisów na polski'
            self.worker = threading.Thread(target=self.translate_worker, daemon=True)
            self.worker.start()

    def translate_worker(self):
        try:
            catalogs = getattr(self, '_translate_catalogs', None)
            records = self.firms_to_translate() if catalogs is None else self.selected_records(catalogs, adopt=True)
            if catalogs is not None:
                self.log('Tłumaczę opisy z katalogów: ' + ', '.join(self.catalog_name(key) for key in catalogs) + '.')
            saved = self.translate_descriptions(records)
            with self.lock:
                self.progress['translations'] = int(self.progress.get('translations') or 0) + 1
                if saved:
                    self.progress['phase'] = f'Przetłumaczono opisy na polski: {saved}'
                elif self.progress.get('phase') == 'Tłumaczenie opisów na polski':
                    self.progress['phase'] = 'Brak nowych niemieckich lub czeskich opisów'
        except Cancelled:
            with self.lock:
                self.progress['phase'] = 'Zatrzymano tłumaczenie — wyniki częściowe'
        except Exception as exc:
            with self.lock:
                self.progress['phase'] = 'Błąd tłumaczenia — wyniki częściowe'
            self.error('Tłumaczenie opisów', exc)
        finally:
            with self.lock:
                self.running = False
                self.checkpoint()

    @staticmethod
    def description_language(record, text):
        haystack = ' '.join(str(record.get(key) or '') for key in ('website', 'source', 'address', 'country')).casefold()
        if domain(record.get('website') or record.get('source') or '').endswith('.cz') or any(word in haystack for word in ('česko', 'czech', 'tschech')):
            return 'cs'
        if domain(record.get('website') or record.get('source') or '').endswith('.de') or any(word in haystack for word in ('deutschland', 'germany', 'niemie')):
            return 'de'
        lowered = ' ' + text.casefold() + ' '
        if any(word in lowered for word in (' společnost ', ' výroba ', ' česk', ' nabízíme ', ' naše ')):
            return 'cs'
        if any(word in lowered for word in (' unternehmen ', ' hersteller ', ' fertigung ', ' unsere ', ' wir ')):
            return 'de'
        return ''

    def start_locations(self):
        with self.lock:
            if self.running:
                raise ValueError('Poczekaj na zakończenie bieżącego zadania')
            if not self.records:
                files = sorted(self.folder.glob('*/wyniki.json'))
                if not files:
                    raise ValueError('Najpierw wyszukaj firmy')
                saved = json.loads(files[-1].read_text(encoding='utf-8'))
                self.records = saved['records']
                self.discoveries = saved.get('discoveries', [])
                self.progress = saved['progress']
            self.run_id = datetime.now().strftime('%Y%m%d_%H%M%S_%f')
            (self.folder / self.run_id).mkdir()
            self.logs = []
            self.stop.clear()
            self.running = True
            self.worker = threading.Thread(target=self.locations_worker, daemon=True)
            self.worker.start()

    def locations_worker(self):
        try:
            self.locate_records()
            self.progress['phase'] = 'Zakończono lokalizowanie — sprawdź dokładność punktów'
        except Cancelled:
            self.progress['phase'] = 'Zatrzymano lokalizowanie — wyniki częściowe'
        except Exception as exc:
            self.progress['phase'] = 'Błąd lokalizowania — wyniki częściowe'
            self.error('Lokalizowanie', exc)
        finally:
            try:
                self.checkpoint()
                export_xlsx(self.snapshot(), self.folder / self.run_id / 'wyniki.xlsx')
                payload = json.dumps(self.snapshot(), ensure_ascii=False).replace('<', '\\u003c')
                template = (Path(__file__).parent / 'static' / 'map.html').read_text(encoding='utf-8')
                (self.folder / self.run_id / 'mapa.html').write_text(template.replace('/*SNAPSHOT*/null', payload), encoding='utf-8')
            except Exception as exc:
                self.error('Eksport', exc)
            with self.lock:
                self.running = False
                self.checkpoint()

    def lookup_site(self, record):
        """Find the company page by its exact name when the catalog has no working site."""
        catalog_name = record.get('catalog_name') or ''
        if not catalog_name and 'europages.' in domain(record.get('source') or ''):
            from catalogs import europages_company
            found = europages_company('', record.get('source') or '')
            catalog_name = found[0] if found else ''
            if catalog_name:
                record['catalog_name'] = catalog_name
        name = re.sub(r'\s+', ' ', catalog_name or record.get('name') or '').strip().replace('"', '')
        if len(name) < 3:
            return ''
        source_host = domain(record.get('source') or '')
        if source_host == 'ares.gov.cz':
            from catalogs import ares_website
            ico = ((record.get('business_ids') or {}).get('ico') or [''])[0]
            self.log('Szukam oficjalnej strony firmy: ' + ' '.join(part for part in (name,ico) if part))
            try:
                site=ares_website(name,ico)
                if site:record['ares_site_verified']=3
                return site
            except SearchBlocked as exc:
                self.log(str(exc));return ''
            except Cancelled:
                raise
            except Exception as exc:
                self.error(name,exc);return ''
        if 'europages.' in source_host or source_host in ('wlw.de', 'industrystock.com', 'industrystock.de', 'industrystock.pl', 'hannovermesse.de'):
            from catalogs import europages_website
            ico = ((record.get('business_ids') or {}).get('ico') or [''])[0]
            place = ico or record.get('locality') or ''
            self.log('Szukam strony firmy: ' + ' '.join(part for part in (name, place) if part))
            try:
                current = domain(record.get('website') or '')
                return europages_website(name, place, (current,) if current else ()) or ''
            except SearchBlocked as exc:
                self.log(str(exc))
                return ''
            except Cancelled:
                raise
            except Exception as exc:
                self.error(name, exc)
                return ''
        ico = ((record.get('business_ids') or {}).get('ico') or [''])[0]
        query = ' '.join(part for part in ('"' + name + '"', ico) if part)
        self.log('DuckDuckGo, dokładna nazwa: ' + query)
        skip = {domain(record.get('source') or ''), domain(record.get('website') or '')}
        skip.discard('')
        try:
            country = record.get('country') if record.get('country') in COUNTRIES else next((code for code,(label,_language) in COUNTRIES.items() if str(record.get('country') or '').casefold() in (code.casefold(),label.casefold())), 'PL')
            hits = search_web(query, country, {'engine': 'duckduckgo', 'stop': self.stop})
        except SearchBlocked as exc:
            self.log(str(exc))
            return ''
        except Cancelled:
            raise
        except Exception as exc:
            self.error(record.get('name') or query, exc)
            return ''
        for hit in hits:
            url = str(hit.get('url') or '').split('#')[0]
            host = domain(url)
            if not url.startswith(('http://', 'https://')) or not host:
                continue
            if is_domain(host, DENIED) or is_domain(host, DIRECTORIES) or 'europages.' in host or is_domain(host, ('duckduckgo.com', 'bing.com', 'yahoo.com')):
                continue
            if any(is_domain(host, (item,)) for item in skip):
                continue
            return url
        return ''

    def enrich(self, record, fetch, searched=False, shallow=False):
        from geolocation import address_rank, find_address
        if not record.get('website'):
            site = self.lookup_site(record)
            searched = True
            if not site:
                self.log(record['name'] + ': wyszukiwarka nie wskazała oficjalnej strony firmy.')
                return
            with self.lock:
                record['website'] = site
            self.log(record['name'] + ': strona firmy — ' + site)
        previous_text = record.get('page_text') or ''
        pending, visited, texts = [record['website']], set(), []
        priorities = {record['website']: -1}
        for _ in range(2 if shallow else self.config.get('pages', 12)):
            if not pending:
                break
            pending.sort(key=lambda item: priorities.get(item, 9))
            url = pending.pop(0)
            if url in visited:
                continue
            visited.add(url)
            try:
                body, kind, final = fetch.page(url)
                if 'pdf' in kind or url.lower().endswith('.pdf'):
                    from pypdf import PdfReader
                    reader = PdfReader(io.BytesIO(body))
                    text = '\n'.join((p.extract_text() or '')[:20000] for p in reader.pages[:15])
                    parsed = dict(text=text, name='', address=find_address(text), lat=None, lon=None, contacts=[dict(person='', role='', email=e.lower(), phone='', source=final, status='Kontakt PDF — bez przypisania do osoby') for e in set(EMAIL.findall(text))], links=[], business_ids=extract_business_ids(text))
                elif 'html' in kind or not kind:
                    parsed = extract(body, final)
                else:
                    continue
            except Cancelled:
                raise
            except Exception as exc:
                self.error('Podstrona ' + domain(url), exc)
                continue
            texts.append((parsed['text'], final))
            with self.lock:
                if len(texts) == 1 and not record.get('catalog_name') and not record['source'].startswith('https://www.openstreetmap.org/'):
                    record['name'] = parsed['name'] or record['name']
                if not record.get('website'):
                    record['website'] = final
                page_address = parsed.get('address') or ''
                if address_rank(page_address) > address_rank(record.get('address')):
                    from geolocation import map_address
                    hint = record.get('country') or ('DE' if domain(record.get('website') or '').endswith('.de') else '')
                    record['address'] = map_address(page_address, hint, record.get('source') or record.get('website') or '')
                if record['lat'] is None and parsed['lat'] is not None:
                    record['lat'], record['lon'], record['geo_source'] = parsed['lat'], parsed['lon'], final
                for contact in parsed['contacts']:
                    key = tuple(contact.get(k, '').lower() for k in ['person', 'email', 'phone'])
                    if not any(tuple(x.get(k, '').lower() for k in ['person', 'email', 'phone']) == key for x in record['contacts']):
                        record['contacts'].append(contact)
                merge_business_ids(record, parsed.get('business_ids', {}))
            if shallow and any(item.get('email') or item.get('phone') for item in record.get('contacts') or []):
                break
            contact_links = []
            for href, label in parsed['links']:
                priority = crawl_link_priority(href, label)
                if shallow and priority not in (0,):
                    continue
                if priority is not None and domain(href) == domain(final) and href not in visited and href not in pending:
                    pending.append(href)
                    priorities[href] = priority
                    if priority <= 1:
                        contact_links.append(href)
                    if shallow:
                        break
            german = domain(final).endswith('.de') or str(record.get('country') or '').casefold() in ('deutschland', 'germany', 'de', 'niemcy')
            if not shallow and len(texts) == 1 and domain(final) and (not contact_links or german):
                root = '{0.scheme}://{0.netloc}/'.format(urlparse(final))
                guesses = (
                    '/contact', '/kontakt', '/impressum', '/contacto', '/contatti',
                    '/iletisim', '/mentions-legales', '/about-us', '/unternehmen',
                    '/team', '/einkauf',
                )
                if german:
                    guesses = (
                        '/impressum', '/de/impressum', '/en/imprint', '/imprint', '/kontakt',
                        '/de/kontakt', '/en/contact', '/unternehmen/impressum', '/meta/impressum',
                    ) + guesses
                for path in guesses:
                    candidate = urljoin(root, path)
                    if candidate not in visited and candidate not in pending:
                        pending.append(candidate)
                        guess_priority = crawl_link_priority(candidate)
                        contact_path = any(part in path for part in ('impressum', 'imprint', 'kontakt', 'contact', 'anschrift'))
                        priorities[candidate] = 0 if german and contact_path else (guess_priority if guess_priority is not None else 4)
        page_text = '\n'.join(text for text, _url in texts)[:50000] or previous_text
        found = find_address(page_text)
        with self.lock:
            if address_rank(found) > address_rank(record.get('address')):
                from geolocation import map_address
                hint = record.get('country') or ('DE' if domain(record.get('website') or '').endswith('.de') else '')
                record['address'] = map_address(found, hint, record.get('source') or record.get('website') or '')
            if page_text:
                record['page_text'] = page_text[:20000]
            elif record.get('website') and searched and not previous_text:
                record['status'] = 'Nie udało się odczytać strony'
        if not page_text and record.get('website') and not searched:
            site = self.lookup_site(record)
            if site:
                with self.lock:
                    record['website'] = site
                    if record.get('status') == 'Nie udało się odczytać strony':
                        record['status'] = 'Oczekuje'
                self.log(record['name'] + ': strona firmy — ' + site)
                self.enrich(record, fetch, searched=True)
            else:
                with self.lock:
                    record['status'] = 'Nie udało się odczytać strony'
        if texts:
            emails = sum(bool(item.get('email')) for item in record.get('contacts') or [])
            phones = sum(bool(item.get('phone')) for item in record.get('contacts') or [])
            parts = [part for part in (record.get('website') or '', f'{emails} e-mail, {phones} tel.' if emails or phones else '', 'adres' if record.get('address') else '') if part]
            self.log('Pobrano dane: ' + (record.get('name') or 'firma') + ' — ' + ', '.join(parts))

def safe_cell(value):
    if isinstance(value, list):
        value = '; '.join(map(str,value))
    if isinstance(value, str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)[:32000]
        # Force literal content, including externally supplied formula-like names.
        if value.lstrip().startswith(('=', '+', '-', '@')):
            value = "'" + value
    return value

def export_xlsx(snapshot, path):
    book = Workbook()
    firms = book.active
    firms.title = 'Firmy'
    firms.append(['ID', 'Firma / tytuł strony', 'WWW', 'Adres', 'Szerokość', 'Długość', 'Źródło współrzędnych', 'Kategorie', 'Dowód dopasowania', 'Status', 'Źródło firmy', 'Sprawdzono UTC', 'Profil poszukiwań', 'Dokładność lokalizacji', 'Dopasowany adres', 'Województwo', 'Grupy', 'Uzasadnienie roli', 'Zakres lokalizacji', 'Katalog', 'Kraj z katalogu', 'Sprawdzenie OSM', 'NACE', 'PKD', 'NIP', 'REGON', 'KRS', 'IČO', 'DIČ', 'Provozovny', 'Opis po polsku', 'Język opisu', 'Opublikowano'])
    contacts = book.create_sheet('Kontakty')
    contacts.append(['ID firmy', 'Firma', 'Osoba / kandydat', 'Rola / kontekst', 'Email', 'Telefon', 'Źródło', 'Pewność przypisania'])
    leads = book.create_sheet('Tropy')
    leads.append(['Tytuł', 'URL', 'Źródło', 'Kategoria', 'Zapytanie', 'Status', 'Sprawdzono UTC'])
    for r in snapshot['records']:
        business = r.get('business_ids') or {}
        row = [safe_cell(r.get(k, '')) for k in ['id', 'name', 'website', 'address', 'lat', 'lon', 'geo_source', 'category', 'evidence', 'status', 'source', 'checked_at', 'profile', 'geo_precision', 'geo_label', 'province', 'groups', 'role_evidence', 'location_scope', 'catalog', 'country', 'osm_check']]
        establishments = '; '.join(' · '.join(part for part in (site.get('name',''), site.get('address',''), 'IČP '+site.get('icp','') if site.get('icp') else '') if part) for site in r.get('establishments') or [])
        firms.append(row + [safe_cell(business.get(k, [])) for k in BUSINESS_FIELDS] + [safe_cell(establishments), safe_cell(r.get('description_pl', '')), safe_cell(r.get('description_language', '')), safe_cell(r.get('published_at', ''))])
        for c in r['contacts']:
            contacts.append([safe_cell(v) for v in [r['id'], r['name']] + [c.get(k, '') for k in ['person', 'role', 'email', 'phone', 'source', 'status']]])
    for d in snapshot['discoveries']:
        leads.append([safe_cell(d.get(k, '')) for k in ['title', 'url', 'source', 'category', 'query', 'status', 'checked_at']])
    for group in GROUPS:
        sheet = book.create_sheet(group)
        sheet.append([c.value for c in firms[1]])
        for i,r in enumerate(snapshot['records'],2):
            if group in r.get('groups',['Do weryfikacji']):
                sheet.append([c.value for c in firms[i]])
    info = book.create_sheet('Informacje')
    for row in [['PROMOT — wyszukiwanie kontrahentów', ''], ['Stan', snapshot['progress']['phase']], ['Utworzono UTC', now()], ['Kontakty', 'Wyłącznie znalezione publicznie. Brak danych = puste pole. Kandydaci na osoby wymagają weryfikacji.'], ['Dopasowanie', 'Ocena słów kluczowych nie potwierdza zapotrzebowania ani kwalifikacji technicznej.'], ['Lokalizacja', 'OSM lub dane strukturalne strony. Photon: geokodowanie adresu, dokładność podana przy firmie. Przybliżenie miejscowości nie oznacza siedziby.'], ['OSM', '© OpenStreetMap contributors — https://www.openstreetmap.org/copyright — ODbL; sprawdź warunki przed udostępnieniem bazy.'], ['LinkedIn / Google Maps', 'Ręczne wyszukiwanie; brak automatycznego pobierania kontaktów.']]:
        info.append(row)
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='153C47')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if isinstance(cell.value, str) and cell.value.startswith(('https://', 'http://')) and '\n' not in cell.value and ' ' not in cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color='176B9A', underline='single')
        for col in sheet.columns:
            width = min(65, max(15, max(len(str(x.value or '')) for x in col[:40]) + 2))
            sheet.column_dimensions[get_column_letter(col[0].column)].width = width
    for row in firms.iter_rows(min_row=2):
        row[4].number_format = row[5].number_format = '0.000000'
    temporary = Path(path).with_suffix('.tmp.xlsx')
    book.save(temporary)
    temporary.replace(path)
