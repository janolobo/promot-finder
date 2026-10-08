import io
import json
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest.mock import patch, Mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import Research, extract, extract_business_ids, export_xlsx, safe_url, coords, Cancelled, Fetcher, crawl_link_priority, internet_query_terms
from openpyxl import load_workbook

class TranslationProcess:
    returncode = 0

    def __init__(self, translation):
        self.translation = translation
        self.stdin = self
        self.stdout = self
        self.stderr = io.StringIO('')
        self._lines = []
        self._index = 0
        self._chunks = []

    def write(self, data):
        self._chunks.append(data)

    def close(self):
        jobs = json.loads(''.join(self._chunks))
        self._lines = [json.dumps(dict(item, translation=self.translation), ensure_ascii=False) for item in jobs]

    def readline(self):
        if self._index >= len(self._lines):
            return ''
        line = self._lines[self._index] + '\n'
        self._index += 1
        return line

    def poll(self):
        return 0

    def wait(self, timeout=None):
        return 0

    def terminate(self):
        self.returncode = -15


HTML = '''<html><head><title>Test Manufacturing</title><script type="application/ld+json">{
"@type":"Organization", "name":"Test Manufacturing", "address":{"streetAddress":"Testowa 2", "addressLocality":"Cieszyn"},
"geo":{"latitude":49.75,"longitude":18.63},
"employee":{"@type":"Person","name":"Anna Testowa","jobTitle":"Purchasing Manager","email":"anna@example.com","telephone":"+48 123 456 789"}}
</script></head><body><h1>Agricultural machinery manufacturer</h1><a href="mailto:office@example.com">Office</a>
<a href="tel:+48111222333">Phone</a><a href="/kontakt">Kontakt</a><script>hidden@example.com</script></body></html>'''

class CoreTests(unittest.TestCase):
    def test_extracts_explicit_business_identifiers_and_codes(self):
        text = '''Company details: NIP: 525-000-77-85, REGON 012345678,
        KRS: 0000123456. Main PKD 2025: 25.62.Z; additional 28.41.Z.
        NACE Rev. 2: C 25.62, 28.41. Phone: +48 123 456 789.'''
        values = extract_business_ids(text)
        self.assertEqual(values['nip'], ['5250007785'])
        self.assertEqual(values['regon'], ['012345678'])
        self.assertEqual(values['krs'], ['0000123456'])
        self.assertEqual(values['pkd'], ['25.62.Z', '28.41.Z'])
        self.assertEqual(values['nace'], ['C 25.62', '28.41'])

    def test_does_not_treat_unlabelled_numbers_as_business_ids(self):
        values = extract_business_ids('Phone +48 525 000 77 85, postal code 00-001, machine 25.62.')
        self.assertTrue(all(not values[key] for key in values))

    def test_description_language_uses_company_country_and_domain(self):
        self.assertEqual(Research.description_language({'website': 'https://firma.de'}, 'Produkte'), 'de')
        self.assertEqual(Research.description_language({'website': 'https://firma.cz'}, 'Výrobky'), 'cs')
        self.assertEqual(Research.description_language({'website': 'https://firma.pl'}, 'Polski opis'), '')

    def test_deep_fill_prioritizes_multilingual_contact_pages(self):
        self.assertEqual(crawl_link_priority('https://firma.example/contatti'), 0)
        self.assertEqual(crawl_link_priority('https://firma.example/einkauf'), 1)
        self.assertIsNone(crawl_link_priority('https://firma.example/privacy'))
        home = '''<a href="/products">Products</a><a href="/team">Team</a>
        <a href="/contatti">Contatti</a>'''
        pages = {
            'https://firma.example': home,
            'https://firma.example/contatti': '<a href="mailto:office@firma.example">E-mail</a>',
            'https://firma.example/team': '<p>Management Team</p>',
            'https://firma.example/products': '<p>Products</p>',
        }
        visited = []
        def page(url):
            visited.append(url)
            return pages[url].encode(), 'text/html', url
        with tempfile.TemporaryDirectory() as tmp, patch('core.Fetcher.page', side_effect=page):
            research = Research(tmp)
            research.records = [dict(id='1', name='Firma', website='https://firma.example', address='', lat=None, lon=None, contacts=[], source='https://catalog.example/firma', status='Oczekuje', osm_text='', groups=['Do weryfikacji'], category='', evidence='')]
            research.config = dict(pages=4, enrich_web=True, categories=[])
            research.start_fill()
            research.worker.join(10)
        self.assertEqual(visited, ['https://firma.example', 'https://firma.example/contatti', 'https://firma.example/team', 'https://firma.example/products'])
        self.assertEqual(research.records[0]['contacts'][0]['email'], 'office@firma.example')

    def test_fill_merges_business_data_from_company_subpages(self):
        pages = {
            'https://firma.example': '<a href="/contact">Contact</a><p>NACE: 25.62</p>',
            'https://firma.example/contact': '<p>NIP 525-000-77-85</p><p>KRS 0000123456</p>',
        }
        with tempfile.TemporaryDirectory() as tmp, patch('core.Fetcher.page', side_effect=lambda url: (pages[url].encode(), 'text/html', url)):
            research = Research(tmp)
            research.records = [dict(id='1', name='Firma', website='https://firma.example', address='', lat=None, lon=None, contacts=[], source='https://firma.example', status='Oczekuje')]
            research.config = dict(pages=2, enrich_web=True, categories=[])
            research.start_fill()
            research.worker.join(10)
        self.assertEqual(research.records[0]['business_ids']['nace'], ['25.62'])
        self.assertEqual(research.records[0]['business_ids']['nip'], ['5250007785'])
        self.assertEqual(research.records[0]['business_ids']['krs'], ['0000123456'])

    def test_fill_searches_all_websites_before_any_mapping(self):
        events=[]
        rows=[dict(id=str(index),name=f'Firma {index}',website='https://example.com',address='',lat=None,lon=None,contacts=[],source='https://catalog.example/'+str(index),status='Oczekuje') for index in (1,2)]
        with tempfile.TemporaryDirectory() as tmp:
            research=Research(tmp);research.records=rows;research.config=dict(pages=1,enrich_web=True,countries=['PL'])
            with patch.object(Research,'enrich',side_effect=lambda record,fetch:events.append('web-'+record['id'])),patch.object(Research,'ensure_local_indexes'),patch.object(Research,'locate_local_addresses',side_effect=lambda records:events.append('address-batch')),patch.object(Research,'place_record',side_effect=lambda record:events.append('map-'+record['id'])):
                research.start_fill();research.worker.join(10)
        self.assertEqual(events,['web-1','web-2','address-batch','map-1','map-2'])

    def test_search_fills_new_firms_before_mapping_them(self):
        events = []
        item = dict(name='Nowa', website='https://nowa.example', source='https://pgm.org.pl/nowa', address='Katowice', country='Polska', province='śląskie', contacts=[], text='Producent zaworów', lat=None, lon=None)
        kept = dict(id='old', name='Stara', website='https://stara.example', source='https://pgm.org.pl/stara', address='Kraków', lat=50.0, lon=19.9, contacts=[], status='Oczekuje', catalog='PGM — członkowie')
        def enrich(record, fetch, searched=False, shallow=False):
            events.append('web-' + record['name'])
            record['address'] = 'Ul. Nowa 1, 40-001 Katowice'
        def place(record):
            events.append('map-' + record['name'])
            record.update(lat=50.26, lon=19.02)
        with tempfile.TemporaryDirectory() as tmp, patch('core.collect', return_value=iter([item])), patch('core.search_web', side_effect=AssertionError('No search engine')), patch.object(Research, 'enrich', side_effect=enrich), patch.object(Research, 'ensure_local_indexes'), patch.object(Research, 'locate_local_addresses'), patch.object(Research, 'place_record', side_effect=place):
            research = Research(tmp)
            research.records = [kept]
            research.start(dict(categories=[], countries=['PL'], sources=['pgm'], enrich_web=True, geocode=True, max_firms=5))
            research.worker.join(15)
        self.assertFalse(research.running)
        self.assertEqual(events, ['web-Nowa', 'map-Nowa'])
        self.assertEqual(research.records[-1]['address'], 'Ul. Nowa 1, 40-001 Katowice')
        self.assertEqual(research.records[-1]['lat'], 50.26)
        self.assertEqual(kept['lat'], 50.0)

    def test_fill_skips_firms_with_name_description_address_phone_and_email(self):
        complete = dict(id='1', name='Kovárna Polák s.r.o.', website='https://polak.example', address='Varnsdorf, Karolíny Světlé 3018', lat=50.0, lon=14.0, contacts=[dict(email='a@polak.example', phone='+420111222333')], source='https://ares.gov.cz/1', status='Oczekuje', osm_text='Kovárna Polák s.r.o. · CZ-NACE: 0162, 14 · Výroba kovových konstrukcí')
        incomplete = dict(id='2', name='Bez maila', website='https://brak.example', address='Praha 1', lat=None, lon=None, contacts=[dict(phone='+420999')], source='https://ares.gov.cz/2', status='Oczekuje', osm_text='Kuźnia i obróbka')
        nace_only = dict(id='3', name='Tylko kody', website='https://kody.example', address='Praha 2', contacts=[dict(email='a@kody.example', phone='+420888')], source='https://ares.gov.cz/3', status='Oczekuje', osm_text='Tylko kody · CZ-NACE: 25, 256')
        self.assertTrue(Research.data_complete(complete))
        self.assertFalse(Research.data_complete(incomplete))
        self.assertFalse(Research.data_complete(nace_only))
        self.assertEqual(Research.visible_description(complete), 'Výroba kovových konstrukcí')
        events = []
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            research.records = [complete, incomplete, nace_only]
            research.config = dict(pages=1, enrich_web=True, countries=['PL'])
            with patch.object(Research, 'enrich', side_effect=lambda record, fetch: events.append('web-' + record['id'])), patch.object(Research, 'ensure_local_indexes'), patch.object(Research, 'locate_local_addresses'), patch.object(Research, 'place_record', side_effect=lambda record: events.append('map-' + record['id'])):
                research.start_fill()
                research.worker.join(10)
        self.assertEqual(events, ['web-2', 'web-3', 'map-2', 'map-3'])
        self.assertEqual(complete['status'], 'Komplet danych')
        self.assertTrue(complete['data_complete'])
        self.assertEqual(complete['website'], 'https://polak.example')
        self.assertEqual(complete['contacts'][0]['email'], 'a@polak.example')
        self.assertNotEqual(incomplete['status'], 'Komplet danych')
        self.assertIn('Komplet danych: 1 firm', '\n'.join(research.logs))

    def test_explicit_person_and_company_contacts_are_separate(self):
        data = extract(HTML, 'https://example.com')
        anna = next(x for x in data['contacts'] if x['person'] == 'Anna Testowa')
        self.assertEqual(anna['email'], 'anna@example.com')
        self.assertEqual(anna['phone'], '+48 123 456 789')
        office = next(x for x in data['contacts'] if x['email'] == 'office@example.com')
        self.assertEqual(office['person'], '')
        self.assertEqual(data['lat'], 49.75)
        self.assertFalse(any(x['email'] == 'hidden@example.com' for x in data['contacts']))

    def test_malformed_and_missing_coordinates(self):
        self.assertEqual(coords('nan', 5), (None, None))
        self.assertEqual(coords(91, 0), (None, None))
        self.assertEqual(coords(0, 0), (0, 0))
        self.assertEqual(extract('<h1>Test</h1>', 'https://example.com')['lat'], None)

    def test_private_urls_denied(self):
        for url in ['file:///etc/passwd', 'http://user:password@example.com', 'http://127.0.0.1', 'http://[::1]', 'http://example.com:8080']:
            with self.assertRaises(ValueError):
                safe_url(url)

    def test_cancel(self):
        event = threading.Event()
        event.set()
        with self.assertRaises(Cancelled):
            Fetcher(event, lambda _:None).check()

    def test_abort_stops_search_without_waiting_for_provider(self):
        from core import search_web
        stop = threading.Event()
        stop.set()
        with self.assertRaises(Cancelled):
            search_web('test', 'PL', {'engine': 'duckduckgo', 'stop': stop})
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            research.running = True
            research.abort()
            self.assertTrue(research.stop.is_set())
            self.assertEqual(research.progress['phase'], 'Zatrzymywanie…')

    def test_worker_exports_and_deduplicates(self):
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            with patch('core.Fetcher.page', return_value=(HTML.encode(), 'text/html', 'https://example.com')):
                research.start(dict(categories=['agriculture'], countries=['PL'], sources=[], seeds='https://example.com\nhttps://example.com/products', pages=2, max_firms=10))
                deadline = time.time() + 10
                while research.running and time.time() < deadline:
                    time.sleep(.02)
                self.assertFalse(research.running)
                state = research.snapshot()
                self.assertEqual(len(state['records']), 1)
                self.assertEqual(state['records'][0]['status'], 'Oczekuje')
                research.start_fill()
                deadline = time.time() + 10
                while research.running and time.time() < deadline:
                    time.sleep(.02)
                self.assertEqual(research.records[0]['status'], 'Oczekuje')
                self.assertFalse(research.records[0].get('category'))
                research.start_analysis()
                deadline = time.time() + 10
                while research.running and time.time() < deadline:
                    time.sleep(.02)
            self.assertFalse(research.running)
            state = research.snapshot()
            self.assertEqual(state['progress']['done'], 1)
            self.assertIn(state['records'][0]['status'], ('Dopasowanie słów — wymaga kwalifikacji', 'Do sprawdzenia'))
            self.assertIn('analysis_match', state['records'][0])
            folder = Path(tmp) / research.run_id
            self.assertTrue((folder / 'mapa.html').exists())
            self.assertTrue((folder / 'wyniki.xlsx').exists())
            self.assertFalse(json.loads((folder/'wyniki.json').read_text())['running'])
            self.assertNotIn('brave_key', (folder/'wyniki.json').read_text())
            book = load_workbook(folder / 'wyniki.xlsx')
            self.assertEqual(book['Firmy']['E2'].value, 49.75)
            self.assertEqual(book['Firmy']['C2'].hyperlink.target, 'https://example.com')

    def test_xlsx_formula_injection(self):
        with tempfile.TemporaryDirectory() as tmp:
            r = dict(id='1', name='=HYPERLINK("bad")', website='https://example.com', contacts=[], lat=None, lon=None)
            state = dict(records=[r], discoveries=[], progress={'phase':'Zakończono'})
            path = Path(tmp)/'test.xlsx'
            export_xlsx(state,path)
            book = load_workbook(path)
            self.assertEqual(book['Firmy']['B2'].data_type, 's')
            self.assertIsNone(book['Firmy']['E2'].value)
            self.assertEqual(book['Firmy'].freeze_panes,'A2')

    def test_require_key_for_search_sources(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                Research(tmp).start(dict(profile='cnc', categories=['automotive'], countries=['PL'], sources=['web'],engine='brave'))

    def test_manual_sources_alone_are_not_a_search(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                Research(tmp).start(dict(profile='cnc', categories=['automotive'], countries=['PL'], sources=['linkedin']))

    def test_stop_saves_partial_results(self):
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            def stop_on_page(url):
                research.stop.set()
                raise Cancelled()
            with patch('core.Fetcher.page', side_effect=stop_on_page):
                research.start(dict(profile='cnc', categories=['automotive'], countries=['PL'], sources=[], seeds='https://example.com', pages=1, max_firms=1))
                research.worker.join(10)
            self.assertFalse(research.running)
            self.assertTrue(research.progress['phase'].startswith('Zatrzymano'))
            self.assertEqual(research.records, [])  # stopped before the seed page was read
            self.assertTrue((Path(tmp)/research.run_id/'wyniki.xlsx').exists())
            self.assertTrue((Path(tmp)/research.run_id/'mapa.html').exists())

    def test_unavailable_page_is_not_successful_match(self):
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            with patch('core.Fetcher.page', side_effect=ValueError('robots.txt nie pozwala')), patch('core.search_web', return_value=[]):
                research.start(dict(profile='cnc', categories=['automotive'], countries=['PL'], sources=[], seeds='https://example.com', pages=1, max_firms=1, enrich_web=False))
                research.worker.join(10)
                self.assertEqual(research.records[0]['status'], 'Oczekuje')
                research.start_fill(dict(enrich_web=True))
                research.worker.join(10)
            self.assertEqual(research.records[0]['status'], 'Nie udało się odczytać strony')
            self.assertEqual(research.progress['errors'], 2)
            self.assertEqual(research.records[0]['contacts'], [])

    def test_osm_and_search_preserve_coordinates_and_source_types(self):
        def response(url, **kwargs):
            r = Mock()
            r.raise_for_status.return_value = None
            if 'overpass' in url:
                r.json.return_value = {'elements':[{'type':'node','id':123,'lat':49.7,'lon':18.6,'tags':{'name':'OSM fixture','description':'agricultural machinery','website':'https://osm-example.com'}}]}
            else:
                r.json.return_value = {'web':{'results':[{'title':'Company fixture','url':'https://example.com'},{'title':'EEN opportunity','url':'https://een.ec.europa.eu/partnering-opportunities/test'}]}}
            return r
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            with patch('core.requests.get', side_effect=response), patch('core.Fetcher.page', return_value=(HTML.encode(),'text/html','https://example.com')):
                research.start(dict(categories=['agriculture'], countries=['PL'], sources=['osm','web'],brave_key='TEST_SECRET',engine='brave',pages=1,max_firms=3))
                research.worker.join(10)
            self.assertFalse(research.running)
            self.assertEqual(len(research.records),2)
            self.assertEqual(research.records[0]['lat'],49.7)
            self.assertEqual(len(research.discoveries),2)
            self.assertGreaterEqual(research.progress['queries_done'],1)
            for f in (Path(tmp)/research.run_id).glob('*'):
                if f.suffix in ['.json','.html','.log']:
                    self.assertNotIn('TEST_SECRET',f.read_text())

    def test_keyless_discovery_and_contacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            with patch('core.search_web', return_value=[{'title':'Firma', 'url':'https://example.com'}]), patch('core.Fetcher.page', return_value=(HTML.encode(),'text/html','https://example.com')):
                research.start(dict(categories=['agriculture'],countries=['PL'],sources=['web'],engine='duckduckgo',pages=1,max_firms=1))
                research.worker.join(10)
                self.assertTrue(research.records)
                research.start_fill()
                research.worker.join(10)
            self.assertFalse(research.running)
            self.assertEqual(len(research.records),1)
            self.assertTrue(research.records[0]['contacts'])
            self.assertIn(research.records[0]['status'], ('Oczekuje', 'Komplet danych'))
            self.assertEqual(research.progress['errors'],0)

    def test_web_search_hits_appear_in_the_internet_list(self):
        from local_osm import PROVINCES
        hits=[{'title':'Kuźnia Lubelska','url':'https://kuznia-lublin.example'},{'title':'Katalog targowy','url':'https://een.ec.europa.eu/partnering-opportunities/x'}]
        queries=[]
        def capture(query, country, config):
            queries.append(query)
            return hits
        with tempfile.TemporaryDirectory() as tmp, patch('core.search_web', side_effect=capture), patch('core.Fetcher.page', return_value=(HTML.encode(),'text/html','https://kuznia-lublin.example')):
            research=Research(tmp)
            research.start(dict(categories=['agriculture'],recipients=['forged'],countries=['PL'],sources=['web'],engine='duckduckgo',pages=1,max_firms=10,enrich_web=False,geocode=False))
            research.worker.join(30)
        self.assertFalse(research.running)
        self.assertGreaterEqual(len(queries), 1)
        self.assertTrue(any('Polska' in query or 'Poland' in query for query in queries))
        for province in PROVINCES:
            self.assertFalse(any(province in query for query in queries))
        self.assertEqual(research.records[0]['catalog'], 'Internet — strony firm')
        self.assertEqual(research.records[0]['name'], 'Kuźnia Lubelska')
        self.assertEqual(research._catalog_key(research.records[0]), 'web')
        rows=research.catalog_rows()
        self.assertEqual([row['catalog_key'] for row in rows], ['web'])
        self.assertEqual(rows[0]['name'], 'Kuźnia Lubelska')
        self.assertTrue(any(item['url']=='https://een.ec.europa.eu/partnering-opportunities/x' for item in research.discoveries))

    def test_internet_query_terms_cover_czechia_without_provinces(self):
        from local_osm import PROVINCES
        queries=internet_query_terms(['agriculture'], [], 'CZ')
        self.assertGreaterEqual(len(queries), 8)
        blob=' '.join(queries)
        self.assertIn('Česko', blob)
        self.assertIn('Czechia', blob)
        self.assertIn('zeměděl', blob)
        for province in PROVINCES:
            self.assertNotIn(province, blob)

    def test_internet_keeps_searching_until_firm_limit(self):
        calls=[]
        def capture(query, country, config):
            n=len(calls)
            calls.append(query)
            return [{'title':f'Firma {n}-{i}','url':f'https://agro-{n}-{i}.example'} for i in range(4)]
        with tempfile.TemporaryDirectory() as tmp, patch('core.search_web', side_effect=capture):
            research=Research(tmp)
            research.start(dict(categories=['agriculture'],countries=['CZ'],sources=['web'],engine='duckduckgo',pages=1,max_firms=12,enrich_web=False,geocode=False))
            research.worker.join(30)
        self.assertFalse(research.running)
        web=[row for row in research.records if row.get('catalog')=='Internet — strony firm']
        self.assertEqual(len(web), 12)
        self.assertGreaterEqual(len(calls), 3)

    def test_internet_search_runs_again_on_the_same_plan(self):
        hits=[{'title':'Kuźnia Nowa','url':'https://kuznia-nowa.example'}]
        queries=[]
        def capture(query, country, config):
            queries.append(query)
            return hits
        with tempfile.TemporaryDirectory() as tmp, patch('core.search_web', side_effect=capture):
            research=Research(tmp)
            config=dict(categories=['agriculture'],countries=['PL'],sources=['web'],engine='duckduckgo',pages=1,max_firms=10,enrich_web=False,geocode=False)
            research.start(config)
            research.worker.join(30)
            first=len(queries)
            research.start(config)
            research.worker.join(30)
        self.assertFalse(research.running)
        self.assertGreater(first, 0)
        self.assertGreater(len(queries), first)
        self.assertEqual(research.cache_counts.get('web'), 1)

    def test_role_is_not_a_person_and_review_author_is_ignored(self):
        html = '''<div>Inżynier Sprzedaży Marcin Testowy <a href="mailto:marcin@example.com">email</a></div>
        <script type="application/ld+json">{"@type":"Review","author":{"@type":"Person","name":"Anonymous Reviewer"}}</script>
        <a href="tel:0123456789">template phone</a>'''
        data = extract(html, 'https://example.com')
        self.assertEqual(data['contacts'][0]['person'], 'Marcin Testowy')
        self.assertFalse(any(c['person']=='Anonymous Reviewer' for c in data['contacts']))
        self.assertFalse(any(c['phone']=='0123456789' for c in data['contacts']))

    def test_permanent_json_cache_survives_restart_and_counts_catalogs(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / 'wyniki'
            research = Research(folder)
            research.records = [dict(name='Cached GmbH', source='https://vdma.eu/#member-1', catalog='VDMA — lista członków', website='https://cached.example', contacts=[], checked_at='2026-10-03T00:00:00+00:00')]
            research.backup_cache(force=True)
            payload = json.loads((Path(tmp) / 'cache' / 'records.json').read_text(encoding='utf-8'))
            self.assertEqual(payload['catalog_counts']['vdma'], 1)
            restored = Research(folder)
            self.assertEqual(restored.records[0]['name'], 'Cached GmbH')
            self.assertEqual(restored.cache_counts['vdma'], 1)
            self.assertEqual(restored.catalog_rows()[0]['catalog_key'], 'vdma')

    def test_trash_removes_one_catalog_and_keeps_the_others(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / 'wyniki'
            research = Research(folder)
            research.run_id = 'session'
            (folder / 'session').mkdir()
            research.records = [
                dict(name='VDMA GmbH', source='https://vdma.eu/#a', catalog='VDMA — lista członków', website='', contacts=[], checked_at='2026-10-06T00:00:00+00:00'),
                dict(name='PGM Sp. z o.o.', source='https://pgm.org.pl/a', catalog='PGM — członkowie', website='', contacts=[], checked_at='2026-10-06T00:00:00+00:00'),
            ]
            research.discoveries = [
                dict(title='VDMA — lista członków', url='https://www.vdma.eu/pl/mitglieder', source='VDMA — lista członków', query=''),
                dict(title='PGM — członkowie', url='https://pgm.org.pl/czlonkowie/', source='PGM — członkowie', query=''),
            ]
            research.progress['search_resume'] = {'completed_catalogs': ['vdma', 'pgm'], 'phrases': ['een:forging', 'Getriebe'], 'queries': [], 'vdma_page': 8}
            research.backup_cache(force=True)
            removed = research.clear_catalog('vdma')
            self.assertEqual(removed, 1)
            self.assertEqual([row['name'] for row in research.records], ['PGM Sp. z o.o.'])
            self.assertEqual(research.cache_counts['vdma'], 0)
            self.assertEqual(research.cache_counts['pgm'], 1)
            self.assertEqual([item['source'] for item in research.discoveries], ['PGM — członkowie'])
            self.assertNotIn('vdma', research.progress['search_resume']['completed_catalogs'])
            self.assertIn('pgm', research.progress['search_resume']['completed_catalogs'])
            self.assertEqual(research.progress['search_resume']['vdma_page'], 1)
            self.assertIn('een:forging', research.progress['search_resume']['phrases'])
            restored = Research(folder)
            self.assertEqual([row['name'] for row in restored.records], ['PGM Sp. z o.o.'])
            self.assertEqual(restored.cache_counts['vdma'], 0)
            with self.assertRaises(ValueError):
                research.clear_catalog('brak')

    def test_trash_removes_internet_firms_and_keeps_catalogs(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / 'wyniki'
            research = Research(folder)
            research.run_id = 'session'
            (folder / 'session').mkdir()
            research.records = [
                dict(name='Kuźnia Sieć', source='https://kuznia.example', catalog='Internet — strony firm', website='https://kuznia.example', contacts=[], checked_at='2026-10-06T00:00:00+00:00'),
                dict(name='PGM Sp. z o.o.', source='https://pgm.org.pl/a', catalog='PGM — członkowie', website='', contacts=[], checked_at='2026-10-06T00:00:00+00:00'),
            ]
            research.discoveries = [
                dict(title='Kuźnia Sieć', url='https://kuznia.example', source='Internet — strony firm', query='odkuwki Polska'),
                dict(title='Targi', url='https://hannovermesse.de/a', source='Katalogi targowe', query='odkuwki Polska (site:hannovermesse.de)'),
            ]
            research.progress['search_resume'] = {'completed_catalogs': [], 'phrases': [], 'queries': ['odkuwki Polska', 'odkuwki Polska (site:hannovermesse.de)'], 'vdma_page': 1}
            research.backup_cache(force=True)
            removed = research.clear_catalog('web')
            self.assertEqual(removed, 1)
            self.assertEqual([row['name'] for row in research.records], ['PGM Sp. z o.o.'])
            self.assertEqual(research.cache_counts.get('web', 0), 0)
            self.assertEqual(research.cache_counts['pgm'], 1)
            self.assertEqual([item['source'] for item in research.discoveries], ['Katalogi targowe'])
            self.assertEqual(research.progress['search_resume']['queries'], ['odkuwki Polska (site:hannovermesse.de)'])
            restored = Research(folder)
            self.assertEqual([row['name'] for row in restored.records], ['PGM Sp. z o.o.'])
            self.assertEqual(restored.cache_counts.get('web', 0), 0)

    def test_restarted_search_keeps_firms_and_continues(self):
        items = [
            dict(name='Firma A', website='', source='https://pgm.org.pl/a', address='', country='Polska', province='śląskie', contacts=[], text='Firma A', lat=None, lon=None),
            dict(name='Firma B', website='', source='https://pgm.org.pl/b', address='', country='Polska', province='śląskie', contacts=[], text='Firma B', lat=None, lon=None),
        ]
        with tempfile.TemporaryDirectory() as tmp, patch('core.collect', side_effect=lambda *args, **kwargs: iter(items)), patch('core.search_web', side_effect=AssertionError('No search engine')):
            research = Research(tmp)
            research.start(dict(categories=[], countries=['PL'], sources=['pgm'], max_firms=1, enrich_web=False))
            research.worker.join(10)
            self.assertEqual([row['name'] for row in research.records], ['Firma A'])
            research.logs.append('00:00:00  stary log')
            research.start(dict(categories=[], countries=['PL'], sources=['pgm'], max_firms=1, enrich_web=False))
            research.worker.join(10)
            self.assertEqual([row['name'] for row in research.records], ['Firma A', 'Firma B'])
            self.assertFalse(any('stary log' in line for line in research.logs))

    def test_start_clears_previous_log(self):
        items = [dict(name='Firma A', website='', source='https://pgm.org.pl/a', address='', country='Polska', province='śląskie', contacts=[], text='Firma A', lat=None, lon=None)]
        with tempfile.TemporaryDirectory() as tmp, patch('core.collect', side_effect=lambda *args, **kwargs: iter(items)), patch('core.search_web', side_effect=AssertionError('No search engine')):
            research = Research(tmp)
            research.logs = ['00:00:00  stary log']
            research.start(dict(categories=[], countries=['PL'], sources=['pgm'], max_firms=1, enrich_web=False))
            research.worker.join(10)
        self.assertFalse(any('stary log' in line for line in research.logs))
        self.assertTrue(research.logs)

    def test_checked_catalogs_limit_fill_analysis_and_translation(self):
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            research.cached_records = [
                dict(name='PGM firma', catalog='PGM — członkowie', source='https://pgm.org.pl/a', website='', osm_text='Hersteller von Ventilen'),
                dict(name='VDMA firma', catalog='VDMA — lista członków', source='https://vdma.eu/b', website='https://werk.de', osm_text='Hersteller von Ventilen'),
            ]
            self.assertEqual([row['name'] for row in research.selected_records(['pgm'])], ['PGM firma'])
            with self.assertRaisesRegex(ValueError, 'katalog'):
                research.take_catalogs({'catalogs': []})
            research.records = []
            process = TranslationProcess('Producent zaworów')
            with patch('core.subprocess.Popen', return_value=process):
                research.start_translate({'catalogs': ['vdma']})
                research.worker.join(10)
            self.assertFalse(research.running)
            self.assertIn('VDMA', '\n'.join(research.logs))
            self.assertNotIn('PGM firma', '\n'.join(research.logs))

    def test_translate_descriptions_saves_polish_text(self):
        with tempfile.TemporaryDirectory() as tmp, patch('core.subprocess.Popen', return_value=TranslationProcess('Polski opis')):
            research = Research(tmp)
            research.records = [dict(id='1', name='Werk', catalog='VDMA — lista członków', source='https://vdma.eu/1', website='https://werk.de', osm_text='Hersteller von Ventilen', contacts=[], checked_at='2026-10-03T00:00:00+00:00')]
            research.cached_records = [dict(research.records[0]), dict(id='2', name='Ceska', catalog='ARES — rejestr', source='https://ares.gov.cz/2', website='https://firma.cz', osm_text='Výroba ventilů', contacts=[])]
            research.start_translate()
            research.worker.join(10)
            self.assertFalse(research.running)
            self.assertEqual(research.records[0]['description_pl'], 'Polski opis')
            self.assertEqual(research.records[0]['osm_text'], 'Hersteller von Ventilen')
            self.assertEqual(research.cached_records[0]['description_pl'], 'Polski opis')
            self.assertEqual(research.cached_records[0]['osm_text'], 'Hersteller von Ventilen')
            self.assertEqual(research.cached_records[1]['description_pl'], 'Polski opis')
            self.assertEqual(research.cached_records[1]['osm_text'], 'Výroba ventilů')
            log = '\n'.join(research.logs)
            self.assertIn('Argos Translate: 1/2 — ', log)
            self.assertIn('Argos Translate: 2/2 — ', log)
            self.assertEqual(research.progress['phase'], 'Przetłumaczono opisy na polski: 2')

    def test_csv_analysis_marks_competition_cooperation_and_clients(self):
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(tmp)
            research.records = [
                dict(id='1', name='Kuźnia Testowa', osm_text='producent odkuwek', page_text='', address='', contacts=[], groups=['Do weryfikacji'], source='https://pgm.org.pl/a', catalog='PGM — członkowie'),
                dict(id='2', name='Warsztat', osm_text='Usługi CNC i frezowanie', page_text='', address='', contacts=[], groups=['Do weryfikacji'], source='https://pgm.org.pl/b', catalog='PGM — członkowie'),
                dict(id='3', name='Montownia', osm_text='Szukamy odbiorcy odkuwek', page_text='', address='', contacts=[], groups=['Do weryfikacji'], source='https://pgm.org.pl/c', catalog='PGM — członkowie'),
            ]
            research.start_analysis({'analysis_kind': 'competition', 'catalogs': ['pgm']})
            research.worker.join(10)
            self.assertFalse(research.running)
            self.assertIn('Konkurencja', research.records[0]['groups'])
            self.assertEqual(research.records[0]['analysis_match'], 'Konkurencja')
            self.assertEqual(research.records[1]['groups'], ['Do weryfikacji'])
            self.assertEqual(research.records[1]['analysis_match'], '')
            research.start_analysis({'analysis_kind': 'cooperation', 'catalogs': ['pgm']})
            research.worker.join(10)
            self.assertIn('Współpraca', research.records[1]['groups'])
            self.assertIn('Konkurencja', research.records[0]['groups'])
            research.start_analysis({'analysis_kind': 'clients', 'catalogs': ['pgm']})
            research.worker.join(10)
            self.assertIn('Klienci', research.records[2]['groups'])
            self.assertEqual(research.records[2]['analysis_match'], 'Klienci')
            self.assertIn('Konkurencja', research.records[0]['analysis_match'])
            with self.assertRaisesRegex(ValueError, 'analizy'):
                research.start_analysis({'analysis_kind': 'unknown', 'catalogs': ['pgm']})
            for record in research.records:
                record['groups'] = ['Do weryfikacji']
                record['analysis_match'] = ''
            research.start_analysis({'analysis_kind': 'all', 'catalogs': ['pgm']})
            research.worker.join(10)
            self.assertFalse(research.running)
            self.assertEqual(research.records[0]['analysis_match'], 'Konkurencja')
            self.assertEqual(research.records[1]['analysis_match'], 'Współpraca')
            self.assertEqual(research.records[2]['analysis_match'], 'Klienci')

    def test_ares_fill_locates_czech_addresses_when_poland_is_selected(self):
        class Index:
            def __init__(self, root, country='PL'):
                self.country = country
                self.pbf = self
                loaded.append(country)
            def exists(self):
                return True
            def ensure(self, stop, log):
                pass
            def locate_addresses(self, records, stop, log):
                if self.country != 'CZ':
                    return 0
                for record in records:
                    if record.get('address_code'):
                        record.update(lat=50.05, lon=14.42, geo_precision='Dokładny adres z lokalnego OSM (Czechy)')
                return 1
        loaded = []
        with tempfile.TemporaryDirectory() as tmp, patch('core.LocalIndex', Index):
            research = Research(tmp)
            research.config['countries'] = ['PL']
            research.records = [dict(name='Kovárna', catalog='ARES — czeski rejestr firm', source='https://ares.gov.cz/ekonomicke-subjekty?ico=21914371', website='', address='Na křivině 1371/1, 14000 Praha 4, Česká republika', country='CZ', address_code='21914371', lat=None, lon=None, contacts=[], business_ids={'ico': ['21914371']})]
            research.start_fill({'catalogs': ['ares'], 'enrich_web': False})
            research.worker.join(10)
            self.assertFalse(research.running)
            self.assertIn('CZ', loaded)
            self.assertEqual(research.records[0]['lat'], 50.05)
            self.assertEqual(research.records[0]['geo_precision'], 'Dokładny adres z lokalnego OSM (Czechy)')

if __name__ == '__main__':
    unittest.main()
