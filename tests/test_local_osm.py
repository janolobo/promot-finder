import sys
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import Mock, patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from local_osm import LocalIndex, PROVINCES, qualify
from core import Research
import osmium


def fixture(root):
    lines=['<osm version="0.6">'];waylines=[];rellines=[]
    for i,name in enumerate(PROVINCES):
        x=14+i*.3;ids=[i*10+j+1 for j in range(4)]
        for nid,(lon,lat) in zip(ids,[(x,49),(x+.25,49),(x+.25,50),(x,50)]):
            lines.append(f'<node id="{nid}" lon="{lon}" lat="{lat}" version="1"/>')
        waylines.append(f'<way id="{i+1}" version="1">'+''.join(f'<nd ref="{n}"/>' for n in ids+[ids[0]])+'</way>')
        rellines.append(f'<relation id="{i+1}" version="1"><member type="way" ref="{i+1}" role="outer"/><tag k="name" v="województwo {name}"/><tag k="ISO3166-2" v="PL-{i:02}"/><tag k="admin_level" v="4"/><tag k="type" v="boundary"/></relation>')
    lines+=['<node id="1001" lat="49.5" lon="14.1" version="1"><tag k="name" v="Kuźnia Testowa"/><tag k="description" v="automotive"/><tag k="man_made" v="works"/></node>', '<node id="1002" lat="49.5" lon="14.4" version="1"><tag k="name" v="Producent maszyn CNC"/><tag k="man_made" v="works"/></node>']
    xml=root/'fixture.osm';xml.write_text('\n'.join(lines+waylines+rellines+['</osm>']))
    with osmium.SimpleWriter(str(root/'poland-latest.osm.pbf')) as w:
        for o in osmium.FileProcessor(str(xml)): w.add(o)


def foreign_fixture(root):
    xml=root/'czech.osm'
    xml.write_text('<osm version="0.6"><node id="2001" lat="49.8" lon="15.5" version="1"><tag k="name" v="Česká Kovárna"/><tag k="description" v="výrobce automotive"/><tag k="man_made" v="works"/></node><node id="2002" lat="49.81" lon="15.51" version="1"><tag k="addr:street" v="Průmyslová"/><tag k="addr:housenumber" v="12"/><tag k="addr:postcode" v="11000"/><tag k="addr:city" v="Praha"/><tag k="ref:ruian:addr" v="998877"/></node></osm>')
    with osmium.SimpleWriter(str(root/'czech-republic-latest.osm.pbf')) as writer:
        for item in osmium.FileProcessor(str(xml)): writer.add(item)


def german_fixture(root):
    xml=root/'germany.osm'
    xml.write_text('<osm version="0.6"><node id="3001" lat="52.52" lon="13.38" version="1"><tag k="name" v="Berliner Schmiede"/><tag k="description" v="Gesenkschmiede automotive"/><tag k="man_made" v="works"/></node><node id="3002" lat="52.53" lon="13.39" version="1"><tag k="addr:street" v="Industriestraße"/><tag k="addr:housenumber" v="7"/><tag k="addr:postcode" v="10115"/><tag k="addr:city" v="Berlin"/></node></osm>')
    with osmium.SimpleWriter(str(root/'germany-latest.osm.pbf')) as writer:
        for item in osmium.FileProcessor(str(xml)): writer.add(item)


class LocalTests(unittest.TestCase):
    def test_build_cache_and_scope_without_network(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);fixture(root)
            idx=LocalIndex(root);idx.ensure(threading.Event(),lambda x:None)
            self.assertEqual(len(idx.data['companies']),2)
            self.assertEqual(idx.province(49.5,14.1),PROVINCES[0])
            self.assertEqual(idx.province(48,14),'')
            r=Research(root/'wyniki')
            with patch('requests.get',side_effect=AssertionError('offline operation')):
                r.start(dict(countries=['PL'],profile='all',categories=['automotive'],sources=['osm'],max_firms=20,enrich_web=False))
                r.worker.join(10)
            self.assertFalse(r.running)
            self.assertEqual(len(r.records),1)
            self.assertEqual(r.records[0]['groups'],['Do weryfikacji'])
            r.start_analysis()
            r.worker.join(10)
            self.assertEqual(r.records[0]['groups'],['Konkurencja'])
            self.assertEqual(r.records[0]['province'],PROVINCES[0])
            self.assertIsNotNone(r.records[0]['lat'])
            self.assertEqual(r.progress['errors'],0)
            from openpyxl import load_workbook
            b=load_workbook(root/'wyniki'/r.run_id/'wyniki.xlsx')
            self.assertEqual(b['Konkurencja'].max_row,2)
            self.assertEqual(b['Odbiorcy odkuwek'].max_row,1)
    def test_qualification(self):
        self.assertEqual(qualify('Firma bez opisu')[0],['Do weryfikacji'])
        self.assertEqual(qualify('Kuźnia Testowa producent odkuwek')[0],['Konkurencja'])
        self.assertIn('Odbiorcy odkuwek',qualify('Producent zaworów przemysłowych')[0])
        self.assertIn('Kooperacja CNC',qualify('Usługi frezowania CNC')[0])
    def test_foreign_country_uses_osm_service(self):
        def response(url, **kwargs):
            self.assertIn('overpass', url)
            self.assertIn('47.2,5.8,55.1,15.1', kwargs.get('params', {}).get('data', ''))
            reply = Mock()
            reply.raise_for_status.return_value = None
            reply.json.return_value = {'elements': [{'type': 'node', 'id': 9, 'lat': 52.5, 'lon': 13.4, 'tags': {'name': 'Werk Berlin', 'description': 'automotive', 'website': 'https://werk.example'}}]}
            return reply
        with tempfile.TemporaryDirectory() as tmp:
            research = Research(Path(tmp) / 'wyniki')
            with patch('core.requests.get', side_effect=response):
                research.start(dict(countries=['DE'], sources=['osm'], max_firms=5, enrich_web=False))
                research.worker.join(10)
            self.assertFalse(research.running)
            self.assertEqual(len(research.records), 1)
            self.assertEqual(research.records[0]['name'], 'Werk Berlin')
            self.assertIn('wyszukiwanie w serwisie — Deutschland', '\n'.join(research.logs))
            self.assertNotIn('Lokalny OSM', '\n'.join(research.logs))
    def test_czech_country_uses_local_extract(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);foreign_fixture(root)
            index=LocalIndex(root,'CZ');index.ensure(threading.Event(),lambda message:None)
            self.assertEqual(index.data['companies'][0]['country'],'CZ')
            records=[dict(country='CZ',address='Průmyslová 12, 11000 Praha, Česká republika',address_code='998877',lat=None)]
            self.assertEqual(index.locate_addresses(records,threading.Event(),lambda message:None),1)
            self.assertAlmostEqual(records[0]['lat'],49.81)
            research=Research(root/'wyniki')
            with patch('core.requests.get',side_effect=AssertionError('offline operation')):
                research.start(dict(countries=['CZ'],sources=['osm'],max_firms=5,enrich_web=False))
                research.worker.join(10)
            self.assertFalse(research.running)
            self.assertEqual(research.records[0]['name'],'Česká Kovárna')
            self.assertIn('Lokalny OSM', '\n'.join(research.logs))
    def test_german_country_uses_local_extract(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);german_fixture(root)
            index=LocalIndex(root,'DE');index.ensure(threading.Event(),lambda message:None)
            self.assertEqual(index.data['companies'][0]['country'],'DE')
            records=[dict(country='Deutschland',address='Industriestraße 7, 10115 Berlin, Deutschland',lat=None)]
            self.assertEqual(index.locate_addresses(records,threading.Event(),lambda message:None),1)
            self.assertAlmostEqual(records[0]['lat'],52.53)
            research=Research(root/'wyniki')
            with patch('core.requests.get',side_effect=AssertionError('offline operation')):
                research.start(dict(countries=['DE'],sources=['osm'],max_firms=5,enrich_web=False))
                research.worker.join(10)
            self.assertFalse(research.running)
            self.assertEqual(research.records[0]['name'],'Berliner Schmiede')
            self.assertIn('Lokalny OSM', '\n'.join(research.logs))
            self.assertNotIn('wyszukiwanie w serwisie — Deutschland', '\n'.join(research.logs))
    def test_start_requires_country_not_province(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError,'kraj'):
                Research(tmp).start({'provinces':[]})
