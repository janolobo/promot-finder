import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import CATEGORIES, PROFILES, Research, company_list, search_term
from openpyxl import load_workbook

PAGE = '<footer>owner@directory.test</footer><script type="application/ld+json">' + json.dumps({'@type':'ItemList','itemListElement':[
 {'@type':'ListItem','item':{'@type':'Organization','name':'Firma A','url':'/firma/a','email':'a@example.com'}},
 {'@type':'ListItem','item':{'@type':'Organization','name':'Firma B','url':'/firma/b','telephone':'+48 222 333 444'}}]}) + '</script>'

class ProfileTests(unittest.TestCase):
    def test_profiles_change_queries(self):
        self.assertEqual(len(CATEGORIES), 5)
        self.assertNotIn('cnc', CATEGORIES)
        for country in ('PL','DE','CZ'):
            self.assertNotIn('odkuwki dostawcy',search_term('automotive',country,'forgings'))
            self.assertNotEqual(search_term('automotive',country,'all'),search_term('automotive',country,'cnc'))

    def test_directory_contacts_are_scoped(self):
        rows = company_list(PAGE, 'https://yoys.pl/adres/test')
        self.assertEqual([r['name'] for r in rows], ['Firma A','Firma B'])
        self.assertEqual(rows[0]['parsed']['contacts'][0]['email'], 'a@example.com')
        self.assertNotIn('owner@directory.test', str(rows))
        self.assertNotIn('a@example.com', str(rows[1]))

    def test_directory_import_and_export(self):
        with tempfile.TemporaryDirectory() as tmp, patch('core.Fetcher.page', return_value=(PAGE, 'text/html', 'https://yoys.pl/adres/test')):
            r = Research(tmp)
            r.start(dict(profile='finished_parts', categories=['automotive'], countries=['PL'], sources=[], seeds='https://yoys.pl/adres/test', max_firms=5))
            r.worker.join(10)
            self.assertFalse(r.running)
            self.assertEqual(len(r.records), 2)
            self.assertEqual(r.records[0]['website'], '')
            self.assertEqual(r.records[1]['source'], 'https://yoys.pl/firma/b')
            self.assertEqual(r.records[0]['profile'], 'Odbiorcy części gotowych')
            book=load_workbook(Path(tmp)/r.run_id/'wyniki.xlsx')
            self.assertEqual(book['Firmy'].max_row,3)
            self.assertEqual(book['Firmy'].cell(2,13).value,'Odbiorcy części gotowych')

    def test_cards_and_navigation(self):
        html='<nav><h2><a href="/about">About</a></h2></nav><div class="company-card"><h2><a href="/a">Firma A</a></h2><a href="mailto:a@example.com">Email</a></div><div class="company-card"><h2><a href="/b">Firma B</a></h2></div>'
        self.assertEqual(len(company_list(html,'https://directory.test/list')), 2)

    def test_invalid_profile(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, 'profil'):
                Research(tmp).start(dict(profile='bogus'))
